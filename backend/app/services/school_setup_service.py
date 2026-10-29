"""School setup in one workbook (#210, 29 Oct 2026; PHASE8_SCHOOL_SETUP_DESIGN.md).

Shailesh, 29 Oct 2026: the school-setup spreadsheet and its loader, so the
first school can start the day Ashalatha's lists come back. One workbook
holds the school's teachers (with the sections and subject each teaches)
and its students (class, section, roll number). It is checked whole before
anything is written, and imported all or nothing: a file with one bad row
creates nobody, unlike the per-role bulk sheet on People, which imports the
rows it can.

Re-uploading a corrected file is safe: people already on Krama are matched
and left as they are (a teacher by email, else mobile, else name in the
school; a student by class, section and roll number, else name in that
section), and a matched teacher still gets any section in the sheet they
don't teach yet.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import School, Student, Teacher, User
from app.models.curriculum import Board, BoardCourse, ClassLevel
from app.models.teacher_assignment import TeacherSectionAssignment

IST = timezone(timedelta(hours=5, minutes=30))
MAX_ROWS = 2000
MAX_SHEET_ROWS = 20000  # blank rows included: reading stops here, so a huge sheet costs little
MAX_COLUMNS = 40
DEFAULT_SUBJECT = "Mathematics"
SUBJECT_ALIASES = {"maths": "mathematics", "math": "mathematics"}
TEACHER_COLUMNS = ["Full Name", "Email", "Mobile", "Subject", "Sections"]
STUDENT_COLUMNS = ["Full Name", "Class", "Section", "Roll Number"]
_SECTION_TOKEN = re.compile(r"^(?:class\s*)?(\d{1,2})\s*[-–\s]?\s*([A-Za-z]{1,3})$", re.IGNORECASE)
_SECTION_NAME = re.compile(r"^[A-Za-z]{1,3}$")
_ROLL = re.compile(r"^[A-Za-z0-9/-]{1,20}$")

# Said the same way whoever holds the address or number (another school's
# teacher, a student, an admin), so the sheet can't be used to find out who
# is registered on the platform (routes_roster.EMAIL_NOT_ACCEPTED_MESSAGE).
EMAIL_NOT_ACCEPTED = "This email address can't be used. Check it is typed correctly, or leave it blank."
MOBILE_NOT_ACCEPTED = "This mobile number can't be used. Check it is typed correctly, or leave it blank."


class SetupFileError(ValueError):
    """The file itself can't be read as a setup workbook."""


# --- reading the workbook --------------------------------------------------------------------


def _key(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _cell(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = " ".join(str(value).split())
    return text or None


def _sheet_rows(wb, name: str, columns: list[str]) -> list[tuple[int, dict]] | None:
    """(sheet row number, {column: text}) for each non-empty row of the sheet
    called `name` (matched by its letters), or None if there is no such sheet."""
    sheet = next((wb[s] for s in wb.sheetnames if _key(s) == _key(name)), None)
    if sheet is None:
        return None
    rows = sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)
    header = next(rows, None) or []
    wanted = {_key(c): c for c in columns}
    index = {}
    for i, heading in enumerate(header):
        column = wanted.get(_key(heading))
        if column and column not in index:
            index[column] = i
    if "Full Name" not in index:
        raise SetupFileError(f"The {name} sheet's first row must be its column headings, including Full Name.")
    out = []
    for number, raw in enumerate(rows, start=2):
        if number > MAX_SHEET_ROWS:
            raise SetupFileError(
                f"The {name} sheet runs past row {MAX_SHEET_ROWS:,}. Delete the empty rows at its end, or split it into two files."
            )
        if raw is None or all(v is None or not str(v).strip() for v in raw):
            continue
        cells = {c: _cell(raw[i]) if i < len(raw) else None for c, i in index.items()}
        if is_example_row(cells):
            continue  # the template's grey example row, if left in
        out.append((number, cells))
        if len(out) > MAX_ROWS:
            raise SetupFileError(f"The {name} sheet has more than {MAX_ROWS} rows. Split it into two files.")
    return out


def read_workbook(raw: bytes) -> tuple[list[tuple[int, dict]], list[tuple[int, dict]]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # any unreadable file gets the same answer
        raise SetupFileError("We couldn't read that file. Upload the School Setup workbook (.xlsx).") from exc
    teachers = _sheet_rows(wb, "Teachers", TEACHER_COLUMNS)
    students = _sheet_rows(wb, "Students", STUDENT_COLUMNS)
    if teachers is None and students is None:
        raise SetupFileError(
            "That file has no Teachers or Students sheet. Download the School Setup template to start from."
        )
    teachers, students = teachers or [], students or []
    if not teachers and not students:
        raise SetupFileError("The Teachers and Students sheets are empty.")
    if len(teachers) + len(students) > MAX_ROWS:
        raise SetupFileError(
            f"This file has {len(teachers) + len(students)} rows, over the {MAX_ROWS}-row limit. Split it into two files."
        )
    return teachers, students


# --- the plan --------------------------------------------------------------------------------


@dataclass
class Problem:
    sheet: str
    row: int
    message: str

    def as_dict(self) -> dict:
        return {"sheet": self.sheet, "row": self.row, "message": self.message}


@dataclass
class TeacherPlan:
    row: int
    full_name: str
    email: str | None
    mobile: str | None
    existing: Teacher | None
    teaches: list[tuple[ClassLevel, str, BoardCourse]] = field(default_factory=list)  # (class, section, course)
    new_teaches: list[tuple[ClassLevel, str, BoardCourse]] = field(default_factory=list)


@dataclass
class StudentPlan:
    row: int
    full_name: str
    class_level: ClassLevel
    section: str
    roll: str | None
    existing: Student | None


@dataclass
class Plan:
    problems: list[Problem] = field(default_factory=list)
    teachers: list[TeacherPlan] = field(default_factory=list)
    students: list[StudentPlan] = field(default_factory=list)

    def summary(self) -> dict:
        new_teachers = [t for t in self.teachers if t.existing is None]
        new_students = [s for s in self.students if s.existing is None]
        sections = sorted({(s.class_level.display_order, s.class_level.code, s.section) for s in self.students})
        return {
            "ok": not self.problems,
            "problems": [p.as_dict() for p in self.problems],
            "newTeachers": len(new_teachers),
            "teachersAlreadyHere": len(self.teachers) - len(new_teachers),
            "newStudents": len(new_students),
            "studentsAlreadyHere": len(self.students) - len(new_students),
            "newSectionsTaught": sum(len(t.new_teaches) for t in self.teachers),
            "sections": [f"{code}{section}" for _o, code, section in sections],
        }


def normal_mobile(text: str) -> str | None:
    digits = re.sub(r"[\s()-]", "", text)
    if digits.startswith("+91"):
        digits = digits[3:]
    elif digits.startswith("91") and len(digits) == 12:
        digits = digits[2:]
    elif digits.startswith("0") and len(digits) == 11:
        digits = digits[1:]
    return digits if re.fullmatch(r"[6-9]\d{9}", digits) else None


def _roll_key(text: str) -> str:
    """A roll number as compared: "01" and 1 (Excel drops the zero) are the same."""
    key = text.strip().upper()
    return str(int(key)) if key.isdigit() else key


def _norm_name(text: str) -> str:
    return " ".join(text.split()).casefold()


def _courses(db: Session, school: School) -> dict[tuple[str, str], list[BoardCourse]]:
    """(class code, subject key) -> the school's board's published courses
    of that name (one, normally)."""
    out: dict[tuple[str, str], list[BoardCourse]] = {}
    if not school.board:
        return out
    rows = (
        db.query(BoardCourse, ClassLevel)
        .join(Board, Board.id == BoardCourse.board_id)
        .join(ClassLevel, ClassLevel.id == BoardCourse.class_level_id)
        .filter(func.upper(Board.code) == school.board.strip().upper(), BoardCourse.status == "PUBLISHED")
        .all()
    )
    for course, level in rows:
        out.setdefault((level.code, _key(course.display_name)), []).append(course)
    return out


def make_plan(
    db: Session,
    school: School,
    teachers: list[tuple[int, dict]],
    students: list[tuple[int, dict]],
    *,
    final: bool = False,
) -> Plan:
    """Every problem in the workbook, or what importing it would do. Reads only.

    Whether a new teacher's email or mobile is free on the rest of Krama is
    asked only when `final` (the import itself), never by the check: a check
    writes nothing and leaves no trace, so answering it there would let any
    admin test thousands of addresses and numbers a minute for who is on the
    platform. An import that is refused leaves an audit row (A14)."""
    plan = Plan()
    levels = {lv.code: lv for lv in db.query(ClassLevel).all()}
    courses = _courses(db, school)
    if not school.board:
        plan.problems.append(
            Problem(
                "Teachers",
                1,
                "The school has no board set (CBSE or ICSE), so its subjects can't be found. Ask the Krama team to set it.",
            )
        )

    school_teachers = (
        db.query(Teacher, User).join(User, User.id == Teacher.user_id).filter(Teacher.school_id == school.id).all()
    )
    by_email = {u.email.lower(): t for t, u in school_teachers if u.email}
    by_phone = {u.phone: t for t, u in school_teachers if u.phone}
    by_tname: dict[str, list[Teacher]] = {}
    for t, u in school_teachers:
        by_tname.setdefault(_norm_name(u.full_name), []).append(t)
    current = {
        (a.class_level_id, a.section.upper(), a.board_course_id): a
        for a in db.query(TeacherSectionAssignment).filter(
            TeacherSectionAssignment.school_id == school.id, TeacherSectionAssignment.end_date.is_(None)
        )
    }
    teacher_names = {t.id: u.full_name for t, u in school_teachers}
    teacher_names_by_id = teacher_names

    seen_emails: dict[str, int] = {}
    seen_mobiles: dict[str, int] = {}
    claimed: dict[tuple, tuple[int, str]] = {}  # (class id, section, course id) -> (row, teacher key) in this file
    teacher_rows: dict[str, TeacherPlan] = {}  # one teacher, several subject rows
    name_idents: dict[str, str] = {}  # name -> the first planned teacher of that name in this file
    for row, cells in teachers:
        name = cells.get("Full Name")
        if not name:
            plan.problems.append(Problem("Teachers", row, "Full Name is empty."))
            continue
        email = (cells.get("Email") or "").strip().lower() or None
        mobile_text = cells.get("Mobile")
        mobile = normal_mobile(mobile_text) if mobile_text else None
        bad = False
        if email and ("@" not in email or email.startswith("@") or email.endswith("@") or " " in email):
            plan.problems.append(Problem("Teachers", row, "Email doesn't look like an email address."))
            bad = True
        if mobile_text and not mobile:
            plan.problems.append(Problem("Teachers", row, "Mobile must be a 10-digit Indian mobile number."))
            bad = True
        # Who this row is: a teacher already at the school, or one named earlier in this file.
        existing = None
        if email and email in by_email:
            existing = by_email[email]
        elif mobile and mobile in by_phone:
            existing = by_phone[mobile]
        elif not email and not mobile and len(by_tname.get(_norm_name(name), [])) == 1:
            existing = by_tname[_norm_name(name)][0]
        elif not email and not mobile and len(by_tname.get(_norm_name(name), [])) > 1:
            plan.problems.append(
                Problem(
                    "Teachers",
                    row,
                    f"There is more than one teacher called {name} on Krama at this school. Add this one's email or mobile.",
                )
            )
            continue
        if existing is not None and _norm_name(teacher_names_by_id[existing.id]) != _norm_name(name):
            # Same school, so naming the teacher is fine; the row and Krama disagree on who this is.
            plan.problems.append(
                Problem(
                    "Teachers",
                    row,
                    f"This email or mobile belongs to {teacher_names_by_id[existing.id]}, who is already on Krama. Check the name or the contact details.",
                )
            )
            continue
        ident = (
            f"id:{existing.id}"
            if existing
            else (f"e:{email}" if email else (f"m:{mobile}" if mobile else f"n:{_norm_name(name)}"))
        )
        planned_name = name_idents.get(_norm_name(name))
        if not existing and not email and not mobile and planned_name:
            ident = planned_name  # a further subject row for a teacher named earlier in the file
        earlier = teacher_rows.get(ident)
        if earlier is not None and _norm_name(earlier.full_name) != _norm_name(name):
            plan.problems.append(
                Problem(
                    "Teachers",
                    row,
                    f"This email or mobile is also on row {earlier.row}, for {earlier.full_name}. Each teacher needs their own, or none.",
                )
            )
            continue
        if (
            planned_name
            and planned_name != ident
            and not teacher_rows[planned_name].email
            and not teacher_rows[planned_name].mobile
        ):
            plan.problems.append(
                Problem(
                    "Teachers",
                    row,
                    f"{name} is also on row {teacher_rows[planned_name].row}, without an email or mobile. Give the same details on each of their rows.",
                )
            )
            continue
        if not existing and final:
            if email and email not in seen_emails and db.query(User.id).filter(User.email == email).first():
                plan.problems.append(Problem("Teachers", row, EMAIL_NOT_ACCEPTED))
                bad = True
            if mobile and mobile not in seen_mobiles and db.query(User.id).filter(User.phone == mobile).first():
                plan.problems.append(Problem("Teachers", row, MOBILE_NOT_ACCEPTED))
                bad = True
        if email:
            first = seen_emails.setdefault(email, row)
            if first != row and teacher_rows.get(ident) is None:
                plan.problems.append(Problem("Teachers", row, f"This email is also on row {first}."))
                bad = True
        if mobile:
            first = seen_mobiles.setdefault(mobile, row)
            if first != row and teacher_rows.get(ident) is None:
                plan.problems.append(Problem("Teachers", row, f"This mobile number is also on row {first}."))
                bad = True
        subject_text = cells.get("Subject") or DEFAULT_SUBJECT
        subject = SUBJECT_ALIASES.get(_key(subject_text), _key(subject_text))
        sections_text = cells.get("Sections") or ""
        tokens = [t.strip() for t in re.split(r"[,;/&]|\band\b", sections_text) if t.strip()]
        if not tokens:
            plan.problems.append(
                Problem("Teachers", row, "Sections is empty: write the sections this teacher teaches, like 5A, 5B.")
            )
            bad = True
        teaches = []
        for token in tokens:
            m = _SECTION_TOKEN.match(token)
            if not m:
                plan.problems.append(
                    Problem("Teachers", row, f'"{token}" isn\'t a section. Write class and section together, like 5A.')
                )
                bad = True
                continue
            level = levels.get(str(int(m.group(1))))
            if level is None:
                plan.problems.append(Problem("Teachers", row, f"Krama has no Class {int(m.group(1))}."))
                bad = True
                continue
            found = courses.get((level.code, subject), []) or [
                c for (code, name), cs in courses.items() if code == level.code and name.endswith(subject) for c in cs
            ]  # "CBSE Class 5 Mathematics" is Mathematics too
            if len(found) != 1:
                if school.board:
                    plan.problems.append(
                        Problem(
                            "Teachers",
                            row,
                            (
                                f"Krama has no {subject_text} for {school.board} Class {level.code}."
                                if not found
                                else f"Krama has more than one {subject_text} course for {school.board} Class {level.code}; ask the Krama team which is the school's."
                            ),
                        )
                    )
                bad = True
                continue
            course = found[0]
            teaches.append((level, m.group(2).upper(), course))
        if bad:
            continue
        tp = teacher_rows.get(ident)
        if tp is None:
            tp = TeacherPlan(row=row, full_name=" ".join(name.split()), email=email, mobile=mobile, existing=existing)
            teacher_rows[ident] = tp
            name_idents.setdefault(_norm_name(name), ident)
            plan.teachers.append(tp)
        for level, section, course in teaches:
            key = (level.id, section, course.id)
            held = current.get(key)
            if held is not None:
                if existing is not None and held.teacher_id == existing.id:
                    continue  # already theirs
                plan.problems.append(
                    Problem(
                        "Teachers",
                        row,
                        f"{level.code}{section} {course.display_name} is already taught by {teacher_names.get(held.teacher_id, 'another teacher')}. "
                        "Moving a class to another teacher is done with Transfer, not in this sheet.",
                    )
                )
                continue
            if key in claimed and claimed[key][1] != ident:
                plan.problems.append(
                    Problem(
                        "Teachers",
                        row,
                        f"{level.code}{section} {course.display_name} is also given to the teacher on row {claimed[key][0]}.",
                    )
                )
                continue
            if key not in claimed:
                claimed[key] = (row, ident)
                tp.teaches.append((level, section, course))
                tp.new_teaches.append((level, section, course))

    # Students
    school_students = (
        db.query(Student, User).join(User, User.id == Student.user_id).filter(Student.school_id == school.id).all()
    )
    by_roll = {}
    by_sname: dict[tuple, list[Student]] = {}
    for s, u in school_students:
        sec = ((s.class_name or "").strip(), (s.section or "").strip().upper())
        if s.roll_number:
            by_roll[(*sec, _roll_key(s.roll_number))] = (s, u)
        by_sname.setdefault((*sec, _norm_name(u.full_name)), []).append(s)
    rolls_in_file: dict[tuple, int] = {}
    names_in_file: dict[tuple, tuple[int, bool]] = {}  # -> (first row, it has a roll number)
    matched: dict[int, int] = {}  # student already on Krama -> the row that is them
    for row, cells in students:
        name = cells.get("Full Name")
        if not name:
            plan.problems.append(Problem("Students", row, "Full Name is empty."))
            continue
        class_text, section_text, roll = cells.get("Class"), cells.get("Section"), cells.get("Roll Number")
        bad = False
        level = None
        if not class_text:
            plan.problems.append(Problem("Students", row, "Class is empty."))
            bad = True
        else:
            m = re.fullmatch(r"(?:class\s*)?(\d{1,2})", class_text, re.IGNORECASE)
            level = levels.get(str(int(m.group(1)))) if m else None
            if level is None:
                plan.problems.append(
                    Problem("Students", row, f'"{class_text}" isn\'t a class Krama has. Write the number, like 5.')
                )
                bad = True
        section = (section_text or "").strip().upper()
        if not section:
            plan.problems.append(Problem("Students", row, "Section is empty."))
            bad = True
        elif not _SECTION_NAME.match(section):
            plan.problems.append(
                Problem("Students", row, f'"{section_text}" isn\'t a section. Write its letter, like A.')
            )
            bad = True
        roll_key = _roll_key(roll) if roll else None
        if roll and not _ROLL.match(roll.strip()):
            plan.problems.append(
                Problem("Students", row, "Roll Number can have letters, digits, / and - only, up to 20.")
            )
            bad = True
        if bad:
            continue
        sec = (level.code, section)
        name_key = (*sec, _norm_name(name))
        existing = None
        if roll_key and (*sec, roll_key) in by_roll:
            s, u = by_roll[(*sec, roll_key)]
            if _norm_name(u.full_name) != _norm_name(name):
                plan.problems.append(
                    Problem(
                        "Students",
                        row,
                        f"Roll number {roll} in {level.code}{section} already belongs to {u.full_name}.",
                    )
                )
                continue
            existing = s
        elif len(by_sname.get(name_key, [])) == 1:
            existing = by_sname[name_key][0]
            if roll_key and existing.roll_number and _roll_key(existing.roll_number) != roll_key:
                plan.problems.append(
                    Problem(
                        "Students",
                        row,
                        f"{name} in {level.code}{section} is already on Krama with roll number {existing.roll_number}.",
                    )
                )
                continue
        elif len(by_sname.get(name_key, [])) > 1:
            plan.problems.append(
                Problem(
                    "Students",
                    row,
                    f"There are already two students called {name} in {level.code}{section}. Add their roll numbers so each can be told apart.",
                )
            )
            continue
        if roll_key:
            first = rolls_in_file.setdefault((*sec, roll_key), row)
            if first != row:
                plan.problems.append(
                    Problem("Students", row, f"Roll number {roll} in {level.code}{section} is also on row {first}.")
                )
                continue
        earlier = names_in_file.setdefault(name_key, (row, bool(roll_key)))
        if earlier[0] != row and not (earlier[1] and roll_key):
            plan.problems.append(
                Problem(
                    "Students",
                    row,
                    f"{name} in {level.code}{section} is also on row {earlier[0]}. If they are two different children, give both a roll number.",
                )
            )
            continue
        if existing is not None:
            first = matched.setdefault(existing.id, row)
            if first != row:
                plan.problems.append(
                    Problem(
                        "Students",
                        row,
                        f"Rows {first} and {row} both match the one {name} in {level.code}{section} already on Krama, "
                        "who has no roll number there, so Krama can't tell which row is that child. Ask the Krama team to add it.",
                    )
                )
                continue
        plan.students.append(
            StudentPlan(
                row=row,
                full_name=" ".join(name.split()),
                class_level=level,
                section=section,
                roll=roll.strip() if roll else None,
                existing=existing,
            )
        )
    plan.problems.sort(key=lambda p: (p.sheet != "Teachers", p.row))
    return plan


# --- importing -------------------------------------------------------------------------------


def apply_plan(db: Session, plan: Plan, *, school: School, creator: User) -> list[dict]:
    """Creates the new people and sections taught, in the caller's
    transaction (it commits or rolls back). Returns the new accounts with
    their first passwords, for the sign-in slips: shown once."""
    from app.api.routes_roster import _create_person

    if plan.problems:
        raise ValueError("The workbook has problems; nothing was imported.")
    today = datetime.now(IST).date()
    created = []
    for tp in plan.teachers:
        teacher = tp.existing
        if teacher is None:
            person = _create_person(
                db,
                creator=creator,
                school=school,
                role="TEACHER",
                full_name=tp.full_name,
                email=tp.email,
                subject_specialization=", ".join(sorted({c.display_name for _l, _s, c in tp.teaches})) or None,
            )
            db.flush()
            user = db.get(User, person["id"])
            if tp.mobile:
                user.phone = tp.mobile
            teacher = db.query(Teacher).filter(Teacher.user_id == user.id).one()
            created.append(
                {
                    "role": "TEACHER",
                    "fullName": tp.full_name,
                    "code": person["code"],
                    "initialPassword": person["initialPassword"],
                    "className": None,
                    "section": None,
                    "rollNumber": None,
                    "teaches": [f"{lv.code}{sec} {c.display_name}" for lv, sec, c in tp.teaches],
                }
            )
        for level, section, course in tp.new_teaches:
            db.add(
                TeacherSectionAssignment(
                    school_id=school.id,
                    teacher_id=teacher.id,
                    class_level_id=level.id,
                    section=section,
                    board_course_id=course.id,
                    start_date=today,
                    created_by_user_id=creator.id,
                )
            )
        db.flush()
    for sp in plan.students:
        if sp.existing is not None:
            if sp.roll and not sp.existing.roll_number:
                sp.existing.roll_number = sp.roll
            continue
        person = _create_person(
            db,
            creator=creator,
            school=school,
            role="STUDENT",
            full_name=sp.full_name,
            email=None,
            class_name=sp.class_level.code,
            section=sp.section,
        )
        db.flush()
        student = db.query(Student).filter(Student.user_id == person["id"]).one()
        student.roll_number = sp.roll
        student.class_level_id = sp.class_level.id
        created.append(
            {
                "role": "STUDENT",
                "fullName": sp.full_name,
                "code": person["code"],
                "initialPassword": person["initialPassword"],
                "className": sp.class_level.code,
                "section": sp.section,
                "rollNumber": sp.roll,
                "teaches": [],
            }
        )
    db.flush()
    return sorted(created, key=slip_order)


def slip_order(slip: dict):
    roll = slip.get("rollNumber") or ""
    number = int(roll) if roll.isdigit() else 10**9
    class_name = slip.get("className") or ""  # People allows any class name: "UKG", "Class 5"
    return (
        slip["role"] != "TEACHER",
        int(class_name) if class_name.isdigit() else 99,
        class_name,
        slip["section"] or "",
        number,
        roll,
        slip["fullName"].casefold(),
    )


# --- the template ----------------------------------------------------------------------------


def template_bytes(school: School) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    how = wb.active
    how.title = "How to Fill In"
    how.column_dimensions["A"].width = 110
    lines = [
        ("t", f"School Setup for {school.name}"),
        (
            "p",
            "Fill in the Teachers and Students sheets, save the file, and upload it on Classes & Sections. Krama checks the whole file first and lists anything to fix; nothing is created until every row is right.",
        ),
        ("h", "Teachers"),
        ("p", "Full Name: as it should appear to students and parents."),
        (
            "p",
            "Email and Mobile: optional. Mobile is a 10-digit Indian number. Each must belong to this teacher alone.",
        ),
        ("p", f"Subject: blank means {DEFAULT_SUBJECT}. A teacher who teaches two subjects takes one row for each."),
        (
            "p",
            "Sections: the sections this teacher teaches that subject in, class and section together, separated by commas: 5A, 5B.",
        ),
        ("h", "Students"),
        (
            "p",
            "Full Name, Class (the number, like 5), Section (its letter, like A) and Roll Number. Nothing else about a child is needed.",
        ),
        ("h", "After the upload"),
        (
            "p",
            "Each new teacher and student gets a code and a first password, shown once, with sign-in slips to print. They choose their own password the first time they sign in.",
        ),
        (
            "p",
            "People already on Krama are recognised and left as they are, so a corrected file can be uploaded again.",
        ),
    ]
    for i, (kind, text) in enumerate(lines, start=1):
        c = how.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Arial", size=14 if kind == "t" else 11, bold=kind in ("t", "h"))
    head_fill = PatternFill("solid", fgColor="FF1F3A5F")
    for title, columns, widths, example in (
        (
            "Teachers",
            TEACHER_COLUMNS,
            [30, 30, 16, 16, 24],
            ["Example: Anita Rao (delete this row)", "anita@example.com", "9876543210", "Mathematics", "5A, 5B"],
        ),
        ("Students", STUDENT_COLUMNS, [34, 8, 10, 14], ["Example: Aarav Sharma (delete this row)", "5", "A", "1"]),
    ):
        ws = wb.create_sheet(title)
        for i, (column, width) in enumerate(zip(columns, widths), start=1):
            c = ws.cell(row=1, column=i, value=column)
            c.font = Font(name="Arial", bold=True, color="FFFFFFFF")
            c.fill = head_fill
            ws.column_dimensions[c.column_letter].width = width
            ex = ws.cell(row=2, column=i, value=example[i - 1])
            ex.font = Font(name="Arial", italic=True, color="FF7F7F7F")
            ex.number_format = "@"
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def is_example_row(cells: dict) -> bool:
    return (cells.get("Full Name") or "").lower().startswith("example:")
