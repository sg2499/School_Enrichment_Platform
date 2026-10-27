"""Loading question patterns from their files, and approving them (#166
step 1, 27 Oct 2026). Used by scripts/load_patterns.py and
scripts/approve_patterns.py; the rules live here so they can be tested.

A pattern file (backend/content/patterns/<engine>/<chapter>.json) names its
engine (board, class, course) and holds the patterns. Loading:

  * checks every pattern (pattern_engine.clean_spec) and tries it 200 times
    (pattern_engine.trial): any failed try, or fewer than MIN_DISTINCT
    different questions, stops the load before anything is written;
  * finds the engine's course, the chapter by number and the skill by its
    Skill ID within that chapter;
  * adds a new pattern as DRAFT, version 1; leaves an unchanged one alone;
    a changed one goes up a version and back to DRAFT (an approval is for
    the spec Ashalatha saw, never for a later one).

Approving reads her reviewed sheet (Pattern ID, Decision, and Version when
the sheet has one): Approve makes the pattern APPROVED, Reject RETIRED,
Change leaves it DRAFT for us to revise. A pattern changed since the sheet
was made is refused.
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from app.services import pattern_engine
from app.services.pattern_engine import PatternError

MIN_DISTINCT = 20
TRIAL_COUNT = 200
AUDIT_LOAD = "content.patterns.loaded"
AUDIT_APPROVE = "content.patterns.reviewed"


class StopBeforeWriting(Exception):
    """Something that must stop the run before it changes anything."""


@dataclass
class LoadPlan:
    add: list = field(default_factory=list)  # dicts ready for QuestionPattern
    change: list = field(default_factory=list)  # (pattern, new fields)
    unchanged: int = 0
    trials: dict = field(default_factory=dict)  # code -> trial result


def _course(db, engine: dict):
    from app.models import Board, BoardCourse, ClassLevel

    board, level, course = (str(engine.get(k) or "").strip() for k in ("board", "classLevel", "course"))
    found = (
        db.query(BoardCourse)
        .join(Board, Board.id == BoardCourse.board_id)
        .join(ClassLevel, ClassLevel.id == BoardCourse.class_level_id)
        .filter(Board.code == board, ClassLevel.code == level, BoardCourse.code == course)
        .all()
    )
    if len(found) != 1:
        raise StopBeforeWriting(f"No single course for {board}, Class {level}, {course} (found {len(found)}).")
    return found[0]


def read_files(paths: list[str]) -> list[tuple[str, dict, dict]]:
    """(file name, engine, raw pattern) for every pattern in the files."""
    out, problems = [], []
    for path in paths:
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            problems.append(f"{path}: can't be read as JSON ({exc}).")
            continue
        engine = data.get("engine") if isinstance(data, dict) else None
        patterns = data.get("patterns") if isinstance(data, dict) else None
        if not isinstance(engine, dict) or not isinstance(patterns, list):
            problems.append(f"{path}: needs an engine and a list of patterns.")
            continue
        for raw in patterns:
            out.append((Path(path).name, engine, raw))
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    return out


def make_plan(
    db, entries: list[tuple[str, dict, dict]], *, trial_count: int = TRIAL_COUNT, min_distinct: int = MIN_DISTINCT
) -> LoadPlan:
    from app.models import Chapter, ConceptLesson
    from app.models.papers import QuestionPattern

    plan, problems, seen = LoadPlan(), [], set()
    courses: dict[str, object] = {}
    for file_name, engine, raw in entries:
        try:
            spec = pattern_engine.clean_spec(raw)
        except PatternError as exc:
            problems.extend(exc.problems)
            continue
        if spec.code in seen:
            problems.append(f"{spec.code}: appears more than once.")
            continue
        seen.add(spec.code)
        key = json.dumps(engine, sort_keys=True)
        if key not in courses:
            try:
                courses[key] = _course(db, engine)
            except StopBeforeWriting as exc:
                problems.append(str(exc))
                courses[key] = None
        course = courses[key]
        if course is None:
            continue
        chapter = (
            db.query(Chapter).filter(Chapter.board_course_id == course.id, Chapter.chapter_no == spec.chapter).all()
        )
        if len(chapter) != 1:
            problems.append(f"{spec.code}: no single Chapter {spec.chapter} in this course.")
            continue
        skill = db.query(ConceptLesson).filter(
            ConceptLesson.chapter_id == chapter[0].id, ConceptLesson.code == spec.skill
        ).all()
        if len(skill) != 1:
            problems.append(f"{spec.code}: no single skill {spec.skill} in Chapter {spec.chapter}.")
            continue
        result = pattern_engine.trial(spec, trial_count)
        plan.trials[spec.code] = result
        if result["failed"]:
            problems.append(f"{spec.code}: {result['failed']} of {result['tried']} tries made no valid question.")
            continue
        if result["distinct"] < min(min_distinct, trial_count):
            problems.append(f"{spec.code}: only {result['distinct']} different questions in {result['tried']} tries.")
            continue
        fields = {
            "code": spec.code, "board_course_id": course.id, "chapter_id": chapter[0].id,
            "concept_lesson_id": skill[0].id, "question_type": spec.question_type, "marks": spec.marks,
            "difficulty": spec.difficulty, "competency": spec.competency,
            "spec_json": json.dumps(raw, ensure_ascii=False, sort_keys=True),
            "spec_sha256": pattern_engine.spec_sha256(raw), "source_file": file_name,
        }
        existing = db.query(QuestionPattern).filter(QuestionPattern.code == spec.code).first()
        if existing is None:
            plan.add.append(fields)
        elif existing.spec_sha256 == fields["spec_sha256"] and existing.concept_lesson_id == fields["concept_lesson_id"]:
            plan.unchanged += 1
        else:
            plan.change.append((existing, fields))
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    return plan


def apply_plan(db, plan: LoadPlan, *, user_id: str | None = None) -> None:
    from app.models.papers import QuestionPattern
    from app.services.audit_service import log_audit_event

    for fields in plan.add:
        db.add(QuestionPattern(**fields, version=1, status="DRAFT"))
    changed = []
    for pattern, fields in plan.change:
        changed.append({"code": pattern.code, "fromVersion": pattern.version, "fromStatus": pattern.status})
        for key, value in fields.items():
            setattr(pattern, key, value)
        pattern.version += 1
        pattern.status = "DRAFT"
        pattern.approved_by = None
        pattern.approved_at = None
    log_audit_event(
        db, AUDIT_LOAD, user_id=user_id,
        details={"added": [f["code"] for f in plan.add], "changed": changed, "unchanged": plan.unchanged},
    )
    db.commit()


# --- approval -------------------------------------------------------------------------

DECISIONS = {"approve": "APPROVED", "reject": "RETIRED", "change": "DRAFT"}


@dataclass(frozen=True)
class Decision:
    row: int
    code: str
    status: str
    version: int | None


def read_decisions(path: str, sheet: str = "Patterns") -> list[Decision]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in book.sheetnames:
        raise StopBeforeWriting(f'The workbook has no "{sheet}" sheet.')
    rows = book[sheet].iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows, ())]
    for need in ("Pattern ID", "Decision"):
        if need not in header:
            raise StopBeforeWriting(f'The "{sheet}" sheet has no column {need}.')
    col = {name: header.index(name) for name in header if name}
    out, problems = [], []
    for number, values in enumerate(rows, start=2):
        cell = lambda name: (values[col[name]] if name in col and col[name] < len(values) else None)  # noqa: E731
        code = str(cell("Pattern ID") or "").strip()
        choice = str(cell("Decision") or "").strip().casefold()
        if not code or not choice:
            continue
        if choice not in DECISIONS:
            problems.append(f"Row {number} ({code}): Decision is {cell('Decision')}, not Approve, Change or Reject.")
            continue
        version = cell("Version")
        if version not in (None, "") and not str(version).strip().isdigit():
            problems.append(f"Row {number} ({code}): Version {version} is not a whole number.")
            continue
        out.append(Decision(number, code, DECISIONS[choice], int(str(version).strip()) if version not in (None, "") else None))
    book.close()
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    return out


def apply_decisions(db, decisions: list[Decision], *, reviewer: str, workbook: str, user_id: str | None = None) -> dict:
    """Records her decisions; refuses all of them if any pattern is missing
    or has changed since the sheet (version, or version 1 when the sheet
    names none)."""
    from app.models.papers import QuestionPattern
    from app.services.audit_service import log_audit_event

    found = {p.code: p for p in db.query(QuestionPattern).filter(
        QuestionPattern.code.in_([d.code for d in decisions])).with_for_update()}
    problems = []
    for d in decisions:
        pattern = found.get(d.code)
        if pattern is None:
            problems.append(f"{d.code}: no such pattern is loaded.")
        elif pattern.version != (d.version or 1):
            problems.append(f"{d.code}: it is now version {pattern.version}, not the version she reviewed.")
    if problems:
        db.rollback()
        raise StopBeforeWriting("\n".join(problems))
    now = datetime.now(timezone.utc)
    changes = []
    for d in decisions:
        pattern = found[d.code]
        if pattern.status == d.status:
            continue
        changes.append({"code": d.code, "version": pattern.version, "from": pattern.status, "to": d.status})
        pattern.status = d.status
        pattern.approved_by = reviewer if d.status == "APPROVED" else None
        pattern.approved_at = now if d.status == "APPROVED" else None
    log_audit_event(db, AUDIT_APPROVE, user_id=user_id,
                    details={"reviewer": reviewer, "workbook": workbook, "changes": changes})
    db.commit()
    return {"changed": len(changes), "unchanged": len(decisions) - len(changes)}
