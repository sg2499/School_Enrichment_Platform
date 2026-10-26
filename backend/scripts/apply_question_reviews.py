"""Records Ashalatha's reviewed Accepted Variants or Marking Guidance (Q2, Q4).

Why this exists (26 Oct 2026). In her answers to the Krama questions she
asked for two review sheets:

  * Q2, "SE_Class_5_Accepted_Variants_Review.xlsx": corrected Accepted
    Variants for Chapters 5 and 15. The workbooks gave many questions a
    stock list ("20|20 mm|20 cm|20 m|20 km|20 markers|20 children|₹20"), so
    the same number in the wrong unit was accepted. Each row proposes the
    list with those entries taken out, keeping the bare number, the unit
    the question asks for and any equal conversion.
  * Q4, "SE_Class_5_Written_Answer_Guides_Review.xlsx": a marking guide
    for each of the 56 three-mark written questions (idea or setup, method
    or working, conclusion with answer and unit), shown to teachers as the
    question's Marking Guidance.

Both sheets have the same shape, so one script records either:

  Question ID | Question | Correct Answer | Current <field> | Proposed <field>
  | Decision | Your <field> | Comments

  * Decision: Approve takes the proposed value, Change takes hers, Keep
    Current leaves the question as it is. A row with no decision is left
    as it is too, so a sheet can come back part-reviewed.

What it does, and why it is safe to run:

  * Checks every row before anything is written: the question exists and
    still has the wording, correct answer and current value she reviewed
    (so a later edit in Curriculum Studio is never overwritten). If any
    row fails, it stops and changes nothing, listing every problem.
  * Without --apply it only reports what would change.
  * With --apply it writes in ONE transaction. For Accepted Variants it
    then marks each changed question's correct answer and every variant
    it keeps; if any would be marked wrong, nothing is saved. It records
    one audit-log entry with the workbook, its SHA-256 and every old and
    new value, so the change can be traced and undone.
  * Running it again is harmless: questions already as she decided are
    left alone.

Usage (from backend/, DATABASE_URL set to the production database):
    python scripts/apply_question_reviews.py --workbook "C:\\path\\to\\SE_Class_5_Accepted_Variants_Reviewed.xlsx"
    python scripts/apply_question_reviews.py --workbook "..." --apply
"""
from __future__ import annotations

import argparse
import hashlib
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Sheet name -> (Question attribute, the field's name in the column headings).
SHEETS = {
    "Accepted Variants": ("accepted_variants", "Accepted Variants"),
    "Marking Guidance": ("teacher_note", "Marking Guidance"),
}
AUDIT_EVENT = "curriculum.question_reviews.applied"
DECISIONS = {"approve": "Approve", "change": "Change", "keep current": "Keep Current"}


class StopBeforeWriting(Exception):
    """Something that must stop the run before it changes anything."""


@dataclass(frozen=True)
class Row:
    row: int
    code: str
    stem: str
    answer: str
    current: str
    decision: str  # Approve | Change | Keep Current
    value: str | None  # what to write; None for Keep Current


@dataclass
class Plan:
    attribute: str
    label: str
    to_change: list = field(default_factory=list)  # (question, Row)
    already_set: int = 0
    kept: int = 0
    undecided: int = 0


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _same(a, b) -> bool:
    """Equal once spacing is ignored; an empty cell, or one showing "—", is empty."""

    def norm(value) -> str:
        text = " ".join(_text(value).split())
        return "" if text in ("—", "-") else text

    return norm(a) == norm(b)


def clean_variants(text: str) -> str:
    """One entry per |, trimmed, no empty or repeated entries."""
    out: list[str] = []
    for part in text.split("|"):
        part = part.strip()
        if part and part not in out:
            out.append(part)
    return "|".join(out)


def read_rows(path: str) -> tuple[str, str, list[Row], int]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [name for name in SHEETS if name in workbook.sheetnames]
    if len(sheets) != 1:
        raise StopBeforeWriting(f"The workbook needs exactly one of the sheets {', '.join(SHEETS)}.")
    sheet = sheets[0]
    attribute, label = SHEETS[sheet]
    columns = {
        "code": "Question ID",
        "stem": "Question",
        "answer": "Correct Answer",
        "current": f"Current {label}",
        "proposed": f"Proposed {label}",
        "decision": "Decision",
        "yours": f"Your {label}",
    }
    rows = workbook[sheet].iter_rows(values_only=True)
    header = [_text(cell) for cell in next(rows, ())]
    missing = [name for name in columns.values() if name not in header]
    if missing:
        raise StopBeforeWriting(f'The "{sheet}" sheet has no column {", ".join(missing)}.')
    index = {key: header.index(name) for key, name in columns.items()}

    out: list[Row] = []
    problems: list[str] = []
    seen: set[str] = set()
    undecided = 0
    for number, values in enumerate(rows, start=2):
        cell = {key: (values[i] if i < len(values) else None) for key, i in index.items()}
        code = _text(cell["code"])
        if not code:
            continue
        where = f"Row {number} ({code})"
        if code in seen:
            problems.append(f"{where}: this question appears more than once.")
            continue
        seen.add(code)
        choice = _text(cell["decision"])
        if not choice:
            undecided += 1
            continue
        decision = DECISIONS.get(choice.casefold())
        if decision is None:
            problems.append(f"{where}: Decision is {choice}, not Approve, Change or Keep Current.")
            continue
        value = None
        if decision == "Approve":
            value = _text(cell["proposed"])
        elif decision == "Change":
            value = _text(cell["yours"])
            if not value:
                problems.append(f"{where}: Decision is Change, but Your {label} is empty.")
                continue
        if value is not None:
            if attribute == "accepted_variants":
                value = clean_variants(value)
            if not value:
                problems.append(f"{where}: the {label} to record is empty.")
                continue
        out.append(Row(number, code, _text(cell["stem"]), _text(cell["answer"]), _text(cell["current"]), decision, value))
    workbook.close()
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    return attribute, label, out, undecided


def make_plan(db, attribute: str, label: str, rows: list[Row], *, lock: bool = False) -> Plan:
    from app.models import Question

    plan = Plan(attribute=attribute, label=label)
    problems: list[str] = []
    query = db.query(Question).filter(Question.code.in_([r.code for r in rows]))
    if lock:
        query = query.with_for_update()
    by_code = {q.code: q for q in query.all()}
    for row in rows:
        question = by_code.get(row.code)
        if question is None:
            problems.append(f"{row.code}: not in the database.")
            continue
        if not _same(question.stem, row.stem) or not _same(question.correct_answer, row.answer):
            problems.append(f"{row.code}: its wording or correct answer has changed since she reviewed it.")
            continue
        now = getattr(question, attribute)
        if row.decision == "Keep Current":
            plan.kept += 1
        elif row.value is not None and _same(now, row.value):
            plan.already_set += 1
        elif not _same(now, row.current):
            problems.append(f"{row.code}: its {label} has been changed since the sheet was made, so it is left for a fresh look.")
        else:
            plan.to_change.append((question, row))
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    return plan


def check_variants_still_right(changed) -> list[str]:
    """After the change: each question's correct answer must still be
    marked right, and every variant it keeps that is a plain amount must
    equal the correct answer (an accepted variant is marked right whatever
    it says, so a slip in her list, "21" for 20, would pass unnoticed)."""
    from app.services import units
    from app.services.learning_service import grade_answer, marking_rule

    failures = []
    for question, row in changed:
        is_correct, _score = grade_answer(question, question.correct_answer)
        if not is_correct:
            failures.append(f'{row.code}: its correct answer "{question.correct_answer}" would be marked wrong.')
        key = units.parse_measurement(question.correct_answer)
        if key is None:
            continue
        # The answer's unit: written on the key, else the question's own
        # (recorded, or read from the question as marking does). "20" on
        # a millimetres question is 20 mm, so "20 cm" is not the same (review).
        key_unit = key.unit or marking_rule(question).answer_unit
        for given in row.value.split("|"):
            amount = units.parse_measurement(given)
            if amount is None:
                continue
            number = amount.number
            if amount.unit and key_unit and amount.unit != key_unit:
                number = units.convert(amount.number, amount.unit, key_unit)
            elif amount.unit and not key_unit:
                number = None  # a unit on an answer that has none
            if number != key.number:
                failures.append(f'{row.code}: "{given}" is not the same amount as the correct answer "{question.correct_answer}".')
    return failures


def apply_plan(db, plan: Plan, *, workbook_name: str, workbook_sha256: str) -> None:
    from app.services.audit_service import log_audit_event

    changes = []
    for question, row in plan.to_change:
        changes.append({"code": row.code, "decision": row.decision, "old": getattr(question, plan.attribute), "new": row.value})
        setattr(question, plan.attribute, row.value)
    db.flush()
    if plan.attribute == "accepted_variants":
        failures = check_variants_still_right(plan.to_change)
        if failures:
            db.rollback()
            raise StopBeforeWriting("Nothing was saved. After the change:\n" + "\n".join(failures))
    log_audit_event(
        db,
        AUDIT_EVENT,
        details={
            "field": plan.attribute,
            "workbook": workbook_name,
            "workbookSha256": workbook_sha256,
            "questionsChanged": len(changes),
            "questionsAlreadySet": plan.already_set,
            "questionsKept": plan.kept,
            "changes": changes,
        },
    )
    db.commit()


def _report(rows: list[Row], plan: Plan, undecided: int) -> None:
    choices = Counter(r.decision for r in rows)
    print(f"Field: {plan.label}")
    print(
        f"Rows decided: {len(rows)} ({choices['Approve']} Approve, {choices['Change']} Change, "
        f"{choices['Keep Current']} Keep Current); no decision yet: {undecided}"
    )
    print(f"Already as she decided: {plan.already_set}")
    print(f"To change: {len(plan.to_change)}")
    for question, row in plan.to_change[:10]:
        print(f"  {row.code}: {row.decision}")
    if len(plan.to_change) > 10:
        print(f"  ... and {len(plan.to_change) - 10} more")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", required=True, help="Ashalatha's reviewed workbook (.xlsx).")
    parser.add_argument("--apply", action="store_true", help="Write the changes. Without it, only report.")
    parser.add_argument("--yes", action="store_true", help="With --apply, skip the typed confirmation.")
    args = parser.parse_args()

    path = Path(args.workbook)
    if not path.is_file():
        print(f"Workbook not found: {path}")
        return 1
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()

    from app.database import SessionLocal, engine
    from scripts.apply_answer_units import mask_database_url

    db_url = str(engine.url)
    print(f"Target database: {mask_database_url(db_url)}")
    if "sqlite" in db_url:
        print("NOTE: DATABASE_URL is not set (or points at SQLite), so this is not production.")
    print(f"Workbook: {path.name} (SHA-256 {sha256[:12]}...)\n")

    db = SessionLocal()
    try:
        attribute, label, rows, undecided = read_rows(str(path))
        plan = make_plan(db, attribute, label, rows)
        _report(rows, plan, undecided)
        if not args.apply:
            print("\nReport only: nothing was written. Run again with --apply to record these.")
            return 0
        if not plan.to_change:
            print("\nNothing to change.")
            return 0
        if not args.yes and input(f"\nType YES to record {len(plan.to_change)} questions' {label}: ").strip() != "YES":
            print("Stopped: nothing was written.")
            return 1
        db.rollback()
        plan = make_plan(db, attribute, label, rows, lock=True)
        apply_plan(db, plan, workbook_name=path.name, workbook_sha256=sha256)
        print(f"\nDone: {len(plan.to_change)} questions updated.")
        return 0
    except StopBeforeWriting as stop:
        print(f"\nStopped before writing anything:\n{stop}")
        return 1
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
