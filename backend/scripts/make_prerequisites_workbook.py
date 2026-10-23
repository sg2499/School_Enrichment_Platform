"""Makes the sheet Ashalatha fills in to say which skills each skill builds on
(#159, 23 Oct 2026), read back by scripts/apply_prerequisites.py.

It lists every skill in the chapter workbooks given (their "Skill Map"
sheets), with the prerequisite each one already describes in words, and
leaves three columns for her: Needs Skill IDs, Minimum Mastery % and
Rescue Note. Nothing touches the database.

Usage (from backend/):
    python scripts/make_prerequisites_workbook.py --workbooks "C:\\...\\Content\\CBSE\\Class-5\\Maths" --out SE_Class_5_Prerequisites.xlsx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

HEADERS = [
    "Chapter No", "Chapter", "Skill ID", "Skill", "Prerequisite Note",
    "Needs Skill IDs", "Minimum Mastery %", "Rescue Note",
]
FOR_HER = {"Needs Skill IDs", "Minimum Mastery %", "Rescue Note"}
HOW_TO = [
    "Which skills does each skill build on?",
    "",
    "For each skill, write in Needs Skill IDs the Skill IDs it builds on, separated by commas,",
    "for example: C5-WT1-S01, C5-WT1-S03. They can be from any chapter.",
    "The Prerequisite Note column is what the chapter workbook already says in words, to help.",
    "",
    "Leave Needs Skill IDs empty for a skill you haven't decided yet: it is left as it is.",
    "Write None for a skill that builds on no other skill.",
    "",
    "Minimum Mastery % (optional): a whole number from 1 to 100. A student at or above it on the",
    "earlier skill counts as secure there. Leave it empty for the platform's default, 70%.",
    "",
    "Rescue Note (optional): a short note for teachers about the rescue practice.",
    "",
    "Foundation Repair uses these links: when a student is weak on a skill and on a skill it builds on,",
    "it recommends practising the earlier one first. Nothing changes on the platform until the sheet",
    "is loaded, and loading it reports every change first.",
]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def read_skills(folder: Path) -> list[list]:
    import openpyxl

    rows: list[list] = []
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "Skill Map" not in book.sheetnames:
            book.close()
            continue
        chapter_no, chapter = None, None
        if "Question Bank" in book.sheetnames:
            bank = book["Question Bank"].iter_rows(values_only=True)
            head = [_text(h) for h in next(bank, ())]
            first = next(bank, None)
            number_col = next((head.index(n) for n in ("Chapter No", "Chapter No.") if n in head), None)
            if first and number_col is not None and "Chapter" in head:
                chapter_no, chapter = first[number_col], _text(first[head.index("Chapter")])
        found_header = False
        for raw in book["Skill Map"].iter_rows(values_only=True):
            cells = [_text(c) for c in raw]
            if not found_header:
                found_header = bool(cells) and cells[0] == "Skill ID"
                continue
            if not cells or not re.match(r"^[A-Z0-9]+-[A-Z0-9]+-S\d+", cells[0] or ""):
                continue
            rows.append([chapter_no, chapter, cells[0], cells[1], cells[3] if len(cells) > 3 else "", "", "", ""])
        book.close()
    rows.sort(key=lambda r: (int(r[0]) if str(r[0] or "").isdigit() else 999, r[2]))
    return rows


def write(rows: list[list], out: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    book = openpyxl.Workbook()
    how = book.active
    how.title = "How To Fill In"
    for i, line in enumerate(HOW_TO, start=1):
        how.cell(row=i, column=1, value=line)
    how["A1"].font = Font(bold=True, size=14)
    how.column_dimensions["A"].width = 110

    sheet = book.create_sheet("Prerequisites")
    sheet.append(HEADERS)
    head_fill = PatternFill("solid", fgColor="2B2370")
    her_fill = PatternFill("solid", fgColor="FFF4D6")
    for col, name in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        sheet.append(row)
    for r in range(2, len(rows) + 2):
        for col, name in enumerate(HEADERS, start=1):
            cell = sheet.cell(row=r, column=col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if name in FOR_HER:
                cell.fill = her_fill
    for col, width in zip("ABCDEFGH", (9, 26, 14, 34, 48, 30, 12, 36)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:H{len(rows) + 1}"
    book.active = 1
    book.save(out)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbooks", required=True, help="Folder holding the chapter workbooks (.xlsx).")
    parser.add_argument("--out", required=True, help="Where to write the sheet (.xlsx).")
    args = parser.parse_args()
    rows = read_skills(Path(args.workbooks))
    if not rows:
        print("No Skill Map rows found in that folder.")
        return 1
    write(rows, Path(args.out))
    print(f"Wrote {len(rows)} skills to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
