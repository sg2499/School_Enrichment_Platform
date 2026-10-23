"""Records Ashalatha's approved answer units on the Class 5 questions (#180).

Why this exists (23 Oct 2026). On 13 Oct she was sent a workbook proposing
an Answer Unit and a Unit Rule for each of the 1,270 Numeric Entry questions
whose Response Format says "unit stated in question" (Chapters 5, 11, 12,
13, 14 and 15). She returned it reviewed: every row Approve or Change, with
a note on each change. This script puts her decisions on those questions.

Why not re-import the chapter workbooks instead: the importer rewrites every
field of every question from the file, its status included, so a re-import
could undo anything changed in Curriculum Studio since and take published
questions back to the file's status. This script changes two fields, Answer
Unit and Unit Rule, on the questions she reviewed, and nothing else. The
other four marking fields (Allow Conversion, Tolerance, Decimal Places,
Unit Mark) are left as they are.

What it does, in order, and why it is safe to run:

  * Reads her decision on every row: Approve takes the proposed unit and
    rule, Change takes hers. It works these out itself rather than trusting
    the workbook's formula columns, which only hold values once Excel has
    calculated them.
  * Checks every row before anything is written: the decision is Approve or
    Change; the unit is one marking knows (the same check Curriculum Studio
    and the importer use); a rule of None has no unit and Stated or Required
    has one; the question exists; and its wording and correct answer are
    still exactly what she reviewed. If any row fails, it stops and changes
    nothing.
  * Without --apply it only reports what would change. Nothing is written.
  * With --apply it writes in ONE transaction, then re-marks every changed
    question's own correct answer, alone and with its unit written after it
    ("24", "24 inches"). If any of those would now be marked wrong, the
    transaction is rolled back untouched. Otherwise it commits, with one
    audit-log entry ("curriculum.answer_units.applied") recording the
    workbook, its SHA-256 and every question's old and new values, so the
    change can be traced and, if ever needed, undone.
  * Running it again is harmless: questions already set as she decided are
    left alone.

Usage (from backend/, with DATABASE_URL set to the production database, as
for scripts/import_class5_maths.py):
    python scripts/apply_answer_units.py --workbook "C:\\path\\to\\SE_Class_5_Answer_Units_Reviewed.xlsx"
    python scripts/apply_answer_units.py --workbook "..." --apply
"""
from __future__ import annotations

import argparse
import hashlib
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

SHEET = "Proposed Units"
COLUMNS = {
    "code": "Question ID",
    "stem": "Question",
    "answer": "Correct Answer",
    "decision": "Decision",
    "proposed_unit": "Proposed Answer Unit",
    "proposed_rule": "Proposed Unit Rule",
    "your_unit": "Your Answer Unit",
    "your_rule": "Your Unit Rule",
    "notes": "Notes",
}
AUDIT_EVENT = "curriculum.answer_units.applied"


class StopBeforeWriting(Exception):
    """Something that must stop the run before it changes anything."""


@dataclass(frozen=True)
class Decision:
    row: int  # the spreadsheet row, for messages
    code: str
    stem: str
    answer: str
    decision: str  # "Approve" or "Change"
    answer_unit: str | None  # a units.UNITS code, or None
    unit_rule: str  # one of learning_service.UNIT_RULES
    notes: str | None


@dataclass
class Plan:
    to_change: list = field(default_factory=list)  # (question, Decision)
    already_set: int = 0
    by_chapter: Counter = field(default_factory=Counter)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def read_decisions(path: str) -> list[Decision]:
    """Her decision on every row, checked. Raises StopBeforeWriting listing
    every row that cannot be used, so they can all be fixed in one go."""
    import openpyxl

    from app.services.learning_service import clean_marking_fields

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise StopBeforeWriting(f'The workbook has no "{SHEET}" sheet.')
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_text(cell) for cell in next(rows, ())]
    missing = [label for label in COLUMNS.values() if label not in header]
    if missing:
        raise StopBeforeWriting(f'The "{SHEET}" sheet has no column {", ".join(missing)}.')
    index = {key: header.index(label) for key, label in COLUMNS.items()}

    decisions: list[Decision] = []
    problems: list[str] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        cell = {key: (row[i] if i < len(row) else None) for key, i in index.items()}
        code = _text(cell["code"])
        if not code:
            continue
        where = f"Row {row_number} ({code})"
        if code in seen:
            problems.append(f"{where}: this question appears more than once.")
            continue
        seen.add(code)
        choice = _text(cell["decision"]).capitalize()
        if choice == "Approve":
            unit_text, rule_text = _text(cell["proposed_unit"]), _text(cell["proposed_rule"])
        elif choice == "Change":
            unit_text, rule_text = _text(cell["your_unit"]), _text(cell["your_rule"])
        else:
            problems.append(f"{where}: Decision is {_text(cell['decision']) or 'empty'}, not Approve or Change.")
            continue
        values, errors = clean_marking_fields({"answer_unit": unit_text, "unit_rule": rule_text})
        if errors:
            problems.extend(f"{where}: {message}" for _key, message in errors)
            continue
        unit, rule = values.get("answer_unit"), values.get("unit_rule")
        if rule is None:
            problems.append(f"{where}: no Unit Rule.")
            continue
        if rule == "NONE" and unit:
            problems.append(f"{where}: Unit Rule None, but an Answer Unit ({unit_text}) is given.")
            continue
        if rule != "NONE" and not unit:
            problems.append(f"{where}: Unit Rule {rule_text}, but no Answer Unit.")
            continue
        decisions.append(
            Decision(
                row=row_number,
                code=code,
                stem=_text(cell["stem"]),
                answer=_text(cell["answer"]),
                decision=choice,
                answer_unit=unit,
                unit_rule=rule,
                notes=_text(cell["notes"]) or None,
            )
        )
    workbook.close()
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    if not decisions:
        raise StopBeforeWriting(f'The "{SHEET}" sheet has no questions in it.')
    return decisions


def _chapter(code: str) -> str:
    parts = code.split("-")
    return parts[2] if len(parts) > 2 else "?"


def make_plan(db, decisions: list[Decision], *, lock: bool = False) -> Plan:
    """What would change. Raises StopBeforeWriting if any question is
    missing or no longer reads as it did when she reviewed it. With
    lock=True the questions are locked until the transaction ends (Postgres
    SELECT ... FOR UPDATE), so nothing can change them between this check
    and the write."""
    from app.models import Question

    plan = Plan()
    problems: list[str] = []
    query = db.query(Question).filter(Question.code.in_([d.code for d in decisions]))
    if lock:
        query = query.with_for_update()
    by_code = {q.code: q for q in query.all()}
    for decision in decisions:
        question = by_code.get(decision.code)
        if question is None:
            problems.append(f"{decision.code}: not in the database.")
            continue
        if _text(question.stem) != decision.stem or _text(question.correct_answer) != decision.answer:
            problems.append(
                f"{decision.code}: its wording or correct answer has changed since she reviewed it, "
                "so her unit may no longer fit."
            )
            continue
        plan.by_chapter[_chapter(decision.code)] += 1
        if (question.answer_unit or None) == decision.answer_unit and (question.unit_rule or "") == decision.unit_rule:
            plan.already_set += 1
        else:
            plan.to_change.append((question, decision))
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    return plan


def check_keys_still_right(changed) -> list[str]:
    """Every changed question's own correct answer must still be marked
    right: as the key is written; and, when the key is a bare number and
    the question has a unit, with that unit after it ("24 inches"). Under
    Required a bare number is meant to lose the unit mark, so there only
    the number with its unit is tried."""
    from app.services import units
    from app.services.learning_service import grade_answer

    failures = []
    for question, decision in changed:
        key = (question.correct_answer or "").strip()
        measured = units.parse_measurement(key)
        key_has_unit = measured is not None and measured.unit is not None
        tries = []
        if key_has_unit or decision.unit_rule != "REQUIRED" or not decision.answer_unit:
            tries.append(key)
        if decision.answer_unit and not key_has_unit:
            tries.append(f"{key} {units.unit_label(decision.answer_unit)}")
        for given in tries:
            is_correct, _score = grade_answer(question, given)
            if not is_correct:
                failures.append(f"{decision.code}: \"{given}\" would be marked wrong.")
    return failures


def apply_plan(db, plan: Plan, *, workbook_name: str, workbook_sha256: str) -> None:
    """Writes the plan in one transaction; rolls back if any answer key
    would be marked wrong afterwards."""
    from app.services.audit_service import log_audit_event

    changes = []
    for question, decision in plan.to_change:
        changes.append(
            {
                "code": decision.code,
                "old": [question.answer_unit, question.unit_rule],
                "new": [decision.answer_unit, decision.unit_rule],
            }
        )
        question.answer_unit = decision.answer_unit
        question.unit_rule = decision.unit_rule
    db.flush()
    failures = check_keys_still_right(plan.to_change)
    if failures:
        db.rollback()
        raise StopBeforeWriting("Nothing was saved. After the change:\n" + "\n".join(failures))
    log_audit_event(
        db,
        AUDIT_EVENT,
        details={
            "workbook": workbook_name,
            "workbookSha256": workbook_sha256,
            "questionsChanged": len(changes),
            "questionsAlreadySet": plan.already_set,
            "changes": changes,
        },
    )
    db.commit()


def _report(decisions: list[Decision], plan: Plan) -> None:
    from app.services import units

    choices = Counter(d.decision for d in decisions)
    print(f"Rows read: {len(decisions)} ({choices['Approve']} Approve, {choices['Change']} Change)")
    print("Questions per chapter: " + ", ".join(f"{ch} {n}" for ch, n in sorted(plan.by_chapter.items())))
    print(f"Already set as she decided: {plan.already_set}")
    print(f"To change: {len(plan.to_change)}")
    rules = Counter(d.unit_rule for _q, d in plan.to_change)
    if rules:
        print("  by Unit Rule: " + ", ".join(f"{rule.title()} {n}" for rule, n in sorted(rules.items())))
    changed_by_her = [d for d in decisions if d.decision == "Change"]
    if changed_by_her:
        print(f"\nHer {len(changed_by_her)} changes:")
        for d in changed_by_her:
            unit = units.unit_label(d.answer_unit) or "no unit"
            print(f"  {d.code}: {unit}, {d.unit_rule.title()}" + (f"  ({d.notes})" if d.notes else ""))


def mask_database_url(url: str) -> str:
    if "@" not in url:
        return url
    scheme_and_creds, host_and_rest = url.rsplit("@", 1)
    scheme = scheme_and_creds.split("://", 1)[0] if "://" in scheme_and_creds else ""
    return f"{scheme}://***:***@{host_and_rest}"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", required=True, help="Ashalatha's reviewed workbook (.xlsx).")
    parser.add_argument("--apply", action="store_true", help="Write the changes. Without it, only report.")
    parser.add_argument("--yes", action="store_true", help="With --apply, skip the typed confirmation.")
    args = parser.parse_args()

    path = Path(args.workbook)
    if not path.is_file():
        print(f"Workbook not found: {path}")
        return 1
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()

    from app.database import SessionLocal, engine

    db_url = str(engine.url)
    print(f"Target database: {mask_database_url(db_url)}")
    if "sqlite" in db_url:
        print("NOTE: DATABASE_URL is not set (or points at SQLite), so this is not production.")
    print(f"Workbook: {path.name} (SHA-256 {sha256[:12]}...)\n")

    db = SessionLocal()
    try:
        decisions = read_decisions(str(path))
        plan = make_plan(db, decisions)
        _report(decisions, plan)
        if not args.apply:
            print("\nReport only: nothing was written. Run again with --apply to record these.")
            return 0
        if not plan.to_change:
            print("\nNothing to change.")
            return 0
        if not args.yes:
            confirm = input(f"\nType YES to record {len(plan.to_change)} questions' units: ")
            if confirm.strip() != "YES":
                print("Stopped: nothing was written.")
                return 1
        # Check again now, with the questions locked, in case one was edited
        # in Curriculum Studio while the prompt was waiting.
        db.rollback()
        plan = make_plan(db, decisions, lock=True)
        apply_plan(db, plan, workbook_name=path.name, workbook_sha256=sha256)
        print(f"\nDone: {len(plan.to_change)} questions updated, each one's correct answer re-marked right.")
        return 0
    except StopBeforeWriting as stop:
        print(f"\nStopped before writing anything:\n{stop}")
        return 1
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
