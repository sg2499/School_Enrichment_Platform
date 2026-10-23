"""Records which skills each skill builds on, from Ashalatha's sheet (#159).

Why this exists (23 Oct 2026). Foundation Repair recommends going back to an
earlier skill when a student is weak on a skill AND on one of the skills it
builds on. Those links (PrerequisiteLink) are empty: the chapter workbooks
only describe prerequisites in words ("Skip counting in 10s, 100s and
1,000s"), which can't be matched to skills reliably by a program. So
Ashalatha fills in a sheet -- one row per skill, with the Skill IDs it needs
-- and this script records them.

The sheet ("Prerequisites", made by scripts/make_prerequisites_workbook.py):

  Skill ID | Skill | Chapter | Prerequisite Note | Needs Skill IDs | Minimum Mastery % | Rescue Note

  * Needs Skill IDs: the Skill IDs this skill builds on, separated by commas
    ("C5-WT1-S01, C5-WT1-S03"). Leave it empty to leave that skill as it is;
    write None to remove every link it has.
  * Minimum Mastery %: optional, a whole number from 1 to 100. A student at
    or above it on the earlier skill is treated as secure there. Empty means
    the platform's default (70%).
  * Rescue Note: optional, a note to teachers about the rescue practice.

What it does, and why it is safe to run:

  * Checks every row before anything is written: each Skill ID names exactly
    one skill on the platform; each needed skill exists and is not the skill
    itself; the mastery figure is a whole number from 1 to 100; and the
    links together go round in no circle (A needs B needs A). If any row
    fails, it stops and changes nothing, listing every problem at once.
  * Without --apply it only reports what would change.
  * With --apply it writes in ONE transaction, with one audit-log entry
    ("curriculum.prerequisites.applied") recording the sheet, its SHA-256
    and every link added, changed and removed.
  * Running it again is harmless: links already as she decided are left
    alone.

Usage (from backend/, DATABASE_URL set to the production database):
    python scripts/apply_prerequisites.py --workbook "C:\\path\\to\\SE_Class_5_Prerequisites_Reviewed.xlsx"
    python scripts/apply_prerequisites.py --workbook "..." --apply
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

SHEET = "Prerequisites"
COLUMNS = {
    "code": "Skill ID",
    "needs": "Needs Skill IDs",
    "mastery": "Minimum Mastery %",
    "note": "Rescue Note",
}
AUDIT_EVENT = "curriculum.prerequisites.applied"
REMOVE_ALL = {"none", "-", "no prerequisites"}


class StopBeforeWriting(Exception):
    """Something that must stop the run before it changes anything."""


@dataclass(frozen=True)
class Row:
    row: int
    code: str
    needs: tuple[str, ...]  # empty with remove_all False: leave as it is
    remove_all: bool
    mastery: int | None
    note: str | None


@dataclass
class Plan:
    # Codes for the report and the audit; ids for writing, since a Skill ID
    # is unique only within its chapter (review).
    add: list[tuple[str, str, int | None, str | None]] = field(default_factory=list)
    change: list[tuple[str, str, tuple, tuple]] = field(default_factory=list)
    # (skill code, needed code, skill id, needed id)
    remove: list[tuple[str, str, str, str]] = field(default_factory=list)
    ids: dict[str, str] = field(default_factory=dict)  # sheet code -> skill id
    unchanged: int = 0
    skills_left_alone: int = 0
    # Needed skills with nothing published to practise: Foundation Repair
    # passes over such a link, so it would do nothing until one is published.
    no_rescue_set: list[str] = field(default_factory=list)

    def summary(self) -> tuple:
        return (sorted(self.add), sorted(self.change), sorted(self.remove))


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def read_rows(path: str) -> list[Row]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise StopBeforeWriting(f'The workbook has no "{SHEET}" sheet.')
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_text(cell) for cell in next(rows, ())]
    missing = [label for label in COLUMNS.values() if label not in header]
    if missing:
        raise StopBeforeWriting(f'The "{SHEET}" sheet has no column {", ".join(missing)}.')
    index = {key: header.index(label) for key, label in COLUMNS.items()}
    out: list[Row] = []
    problems: list[str] = []
    seen: set[str] = set()
    for number, raw in enumerate(rows, start=2):
        cell = {key: (raw[i] if i < len(raw) else None) for key, i in index.items()}
        code = _text(cell["code"])
        if not code:
            continue
        where = f"Row {number} ({code})"
        if code in seen:
            problems.append(f"{where}: this skill appears more than once.")
            continue
        seen.add(code)
        needs_text = _text(cell["needs"])
        remove_all = needs_text.casefold() in REMOVE_ALL
        needs = () if remove_all else tuple(dict.fromkeys(p.strip() for p in re.split(r"[,;\n]", needs_text) if p.strip()))
        mastery_text = _text(cell["mastery"])
        mastery = None
        if mastery_text:
            # A cell Excel formats as a percentage holds 0.7 for 70% (review).
            raw_value = cell["mastery"]
            if isinstance(raw_value, float) and 0 < raw_value <= 1:
                mastery_text = format(round(raw_value * 100, 6), "g")
            mastery_text = mastery_text.rstrip("%").strip()
            try:
                mastery = int(float(mastery_text))
            except ValueError:
                mastery = -1
            if mastery < 1 or mastery > 100 or str(mastery) not in (mastery_text, mastery_text.rstrip("0").rstrip(".")):
                problems.append(f"{where}: Minimum Mastery % must be a whole number from 1 to 100, not {mastery_text}.")
                continue
        note = _text(cell["note"]) or None
        if note and len(note) > 1000:
            problems.append(f"{where}: the Rescue Note is longer than 1,000 characters.")
            continue
        if code in needs:
            problems.append(f"{where}: a skill can't need itself.")
            continue
        out.append(Row(number, code, needs, remove_all, mastery, note))
    workbook.close()
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    if not out:
        raise StopBeforeWriting(f'The "{SHEET}" sheet has no skills in it.')
    return out


def _skills_by_code(db, codes: set[str]) -> tuple[dict, list[str]]:
    from app.models import ConceptLesson

    found: dict[str, list] = {}
    for lesson in db.query(ConceptLesson).filter(ConceptLesson.code.in_(sorted(codes))).all():
        found.setdefault(lesson.code, []).append(lesson)
    problems = []
    for code in sorted(codes):
        if code not in found:
            problems.append(f"{code}: no skill on the platform has this Skill ID.")
        elif len(found[code]) > 1:
            problems.append(f"{code}: more than one skill has this Skill ID, so it can't be matched safely.")
    return {code: lessons[0] for code, lessons in found.items() if len(lessons) == 1}, problems


def _find_circle(edges: dict[str, set[str]]) -> list[str] | None:
    """A circle of needs (A needs B needs A), or None."""
    state: dict[str, int] = {}
    path: list[str] = []

    def visit(node: str) -> list[str] | None:
        state[node] = 1
        path.append(node)
        for nxt in sorted(edges.get(node, ())):
            if state.get(nxt) == 1:
                return path[path.index(nxt):] + [nxt]
            if state.get(nxt) is None:
                found = visit(nxt)
                if found:
                    return found
        path.pop()
        state[node] = 2
        return None

    for node in sorted(edges):
        if state.get(node) is None:
            found = visit(node)
            if found:
                return found
    return None


def make_plan(db, rows: list[Row]) -> Plan:
    from app.models import ConceptLesson, LearningActivity, PrerequisiteLink

    codes = {r.code for r in rows} | {n for r in rows for n in r.needs}
    skills, problems = _skills_by_code(db, codes)
    if problems:
        raise StopBeforeWriting("\n".join(problems))
    plan = Plan(ids={code: lesson.id for code, lesson in skills.items()})
    code_of = {lesson.id: code for code, lesson in skills.items()}
    existing = db.query(PrerequisiteLink).all()
    for lesson in db.query(ConceptLesson).filter(
        ConceptLesson.id.in_(sorted({l.concept_lesson_id for l in existing} | {l.prerequisite_concept_lesson_id for l in existing}))
    ):
        code_of.setdefault(lesson.id, lesson.code)
    # Every skill's needs after the change, by id, to check for a circle.
    after: dict[str, set[str]] = {}
    for link in existing:
        after.setdefault(link.concept_lesson_id, set()).add(link.prerequisite_concept_lesson_id)
    current = {(l.concept_lesson_id, l.prerequisite_concept_lesson_id): l for l in existing}
    for row in rows:
        if not row.needs and not row.remove_all:
            plan.skills_left_alone += 1
            continue
        skill_id = plan.ids[row.code]
        wanted = {plan.ids[n] for n in row.needs}
        after[skill_id] = wanted
        for skill, needed in current:
            if skill == skill_id and needed not in wanted:
                plan.remove.append((row.code, code_of[needed], skill, needed))
        for needed in row.needs:
            link = current.get((skill_id, plan.ids[needed]))
            if link is None:
                plan.add.append((row.code, needed, row.mastery, row.note))
            elif (link.minimum_mastery, link.rescue_activity_note) != (row.mastery, row.note):
                plan.change.append(
                    (row.code, needed, (link.minimum_mastery, link.rescue_activity_note), (row.mastery, row.note))
                )
            else:
                plan.unchanged += 1
    circle = _find_circle(after)
    if circle:
        raise StopBeforeWriting(
            "These skills need each other in a circle: " + " needs ".join(code_of.get(i, i) for i in circle) + "."
        )
    needed_ids = {plan.ids[n] for _s, n, *_ in plan.add}
    if needed_ids:
        practisable = {
            lid
            for (lid,) in db.query(LearningActivity.concept_lesson_id).filter(
                LearningActivity.concept_lesson_id.in_(sorted(needed_ids)),
                LearningActivity.status == "PUBLISHED",
                LearningActivity.activity_type.in_(("REMEDIATION", "CORE_PRACTICE")),
            )
        }
        plan.no_rescue_set = sorted(code_of[i] for i in needed_ids - practisable)
    return plan


def apply_plan(db, plan: Plan, *, workbook_name: str, workbook_sha256: str) -> None:
    from datetime import UTC, datetime, timedelta

    from app.models import PrerequisiteLink
    from app.services.audit_service import log_audit_event

    ids = plan.ids
    for _skill, _needed, skill_id, needed_id in plan.remove:
        db.query(PrerequisiteLink).filter(
            PrerequisiteLink.concept_lesson_id == skill_id,
            PrerequisiteLink.prerequisite_concept_lesson_id == needed_id,
        ).delete(synchronize_session=False)
    for skill, needed, _old, (mastery, note) in plan.change:
        link = (
            db.query(PrerequisiteLink)
            .filter(
                PrerequisiteLink.concept_lesson_id == ids[skill],
                PrerequisiteLink.prerequisite_concept_lesson_id == ids[needed],
            )
            .one()
        )
        link.minimum_mastery, link.rescue_activity_note = mastery, note
    # Foundation Repair tries a skill's links oldest first; within one run
    # they keep the sheet's order, a microsecond apart (review: one
    # transaction's now() would tie them all).
    start = datetime.now(UTC)
    for index, (skill, needed, mastery, note) in enumerate(plan.add):
        db.add(
            PrerequisiteLink(
                concept_lesson_id=ids[skill],
                prerequisite_concept_lesson_id=ids[needed],
                minimum_mastery=mastery,
                rescue_activity_note=note,
                created_at=start + timedelta(microseconds=index),
            )
        )
    log_audit_event(
        db,
        AUDIT_EVENT,
        details={
            "workbook": workbook_name,
            "workbookSha256": workbook_sha256,
            "added": [{"skill": s, "needs": n, "minimumMastery": m, "note": t} for s, n, m, t in plan.add],
            "changed": [{"skill": s, "needs": n, "old": list(o), "new": list(w)} for s, n, o, w in plan.change],
            "removed": [{"skill": s, "needs": n} for s, n, _si, _ni in plan.remove],
        },
    )
    db.commit()


def _report(rows: list[Row], plan: Plan) -> None:
    print(f"Skills in the sheet: {len(rows)} ({plan.skills_left_alone} left empty, so left as they are)")
    print(f"Links to add: {len(plan.add)}")
    print(f"Links to change (mastery or note): {len(plan.change)}")
    print(f"Links to remove: {len(plan.remove)}")
    print(f"Links already as she decided: {plan.unchanged}")
    for skill, needed, mastery, _note in plan.add[:20]:
        print(f"  + {skill} needs {needed}" + (f" (secure at {mastery}%)" if mastery else ""))
    if len(plan.add) > 20:
        print(f"  ... and {len(plan.add) - 20} more")
    for skill, needed, _si, _ni in plan.remove:
        print(f"  - {skill} no longer needs {needed}")
    if plan.no_rescue_set:
        print(
            f"\nNote: {len(plan.no_rescue_set)} needed skills have no published practice set to send a student back to, "
            "so Foundation Repair passes over those links until one is published: " + ", ".join(plan.no_rescue_set)
        )


def mask_database_url(url: str) -> str:
    if "@" not in url:
        return url
    scheme_and_creds, host_and_rest = url.rsplit("@", 1)
    scheme = scheme_and_creds.split("://", 1)[0] if "://" in scheme_and_creds else ""
    return f"{scheme}://***:***@{host_and_rest}"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", required=True, help="Ashalatha's filled-in prerequisites sheet (.xlsx).")
    parser.add_argument("--apply", action="store_true", help="Write the changes. Without it, only report.")
    parser.add_argument("--yes", action="store_true", help="With --apply, skip the typed confirmation.")
    args = parser.parse_args()
    path = Path(args.workbook)
    if not path.is_file():
        print(f"Workbook not found: {path}")
        return 1
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()

    from app.database import SessionLocal, engine

    db_url = str(engine.url)
    print(f"Target database: {mask_database_url(db_url)}")
    if "sqlite" in db_url:
        print("NOTE: DATABASE_URL is not set (or points at SQLite), so this is not production.")
    print(f"Workbook: {path.name} (SHA-256 {sha256[:12]}...)\n")
    db = SessionLocal()
    try:
        rows = read_rows(str(path))
        plan = make_plan(db, rows)
        _report(rows, plan)
        if not args.apply:
            print("\nReport only: nothing was written. Run again with --apply to record these.")
            return 0
        if not (plan.add or plan.change or plan.remove):
            print("\nNothing to change.")
            return 0
        if not args.yes:
            confirm = input("\nType YES to record these links: ")
            if confirm.strip() != "YES":
                print("Stopped: nothing was written.")
                return 1
        db.rollback()
        again = make_plan(db, rows)
        if again.summary() != plan.summary():
            print("\nStopped: the links changed while you were confirming. Nothing was written; run it again to see the new report.")
            return 1
        apply_plan(db, again, workbook_name=path.name, workbook_sha256=sha256)
        print(f"\nDone: {len(plan.add)} added, {len(plan.change)} changed, {len(plan.remove)} removed.")
        return 0
    except StopBeforeWriting as stop:
        print(f"\nStopped before writing anything:\n{stop}")
        return 1
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
