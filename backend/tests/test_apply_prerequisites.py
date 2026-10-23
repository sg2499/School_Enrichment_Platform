"""scripts/apply_prerequisites.py and make_prerequisites_workbook.py (#159,
23 Oct 2026): which skills each skill builds on, from Ashalatha's sheet.

Runs the real script functions against a throwaway in-memory database. The
rules under test: an empty row leaves a skill as it is and None removes its
links; unknown, ambiguous or self-referring skills, a bad mastery figure or
a circle of needs stop the run before anything is written; a second run
changes nothing; one audit entry records every change. Also: the sheet made
from chapter workbooks lists their skills and is read back as it is.
"""
import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import AuditLog, ConceptLesson, PrerequisiteLink


def _load(name):
    path = Path(__file__).resolve().parents[1] / "scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


apply_prerequisites = _load("apply_prerequisites")
make_workbook = _load("make_prerequisites_workbook")
StopBeforeWriting = apply_prerequisites.StopBeforeWriting
HEADERS = make_workbook.HEADERS
CODES = ["C5-WT1-S01", "C5-WT1-S02", "C5-WT1-S03", "C5-FR-S01"]


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, future=True)()
    for i, code in enumerate(CODES, start=1):
        session.add(ConceptLesson(chapter_id="ch-1" if "WT1" in code else "ch-2", code=code, title=f"Skill {i}", sequence=i))
    session.commit()
    yield session
    session.close()


def _sheet(tmp_path, rows, name="filled.xlsx") -> str:
    """rows: (skill, needs, mastery, note)"""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Prerequisites"
    sheet.append(HEADERS)
    for code, needs, mastery, note in rows:
        line = dict.fromkeys(HEADERS)
        line.update({"Skill ID": code, "Needs Skill IDs": needs, "Minimum Mastery %": mastery, "Rescue Note": note})
        sheet.append([line[h] for h in HEADERS])
    path = tmp_path / name
    book.save(path)
    return str(path)


def _links(db):
    by_id = {l.id: l.code for l in db.query(ConceptLesson).all()}
    return sorted(
        (by_id[l.concept_lesson_id], by_id[l.prerequisite_concept_lesson_id], l.minimum_mastery)
        for l in db.query(PrerequisiteLink).all()
    )


def _run(db, path):
    rows = apply_prerequisites.read_rows(path)
    plan = apply_prerequisites.make_plan(db, rows)
    apply_prerequisites.apply_plan(db, plan, workbook_name="filled.xlsx", workbook_sha256="abc")
    return plan


def test_links_are_added_changed_and_removed_and_a_rerun_changes_nothing(db, tmp_path):
    path = _sheet(tmp_path, [
        ("C5-WT1-S02", "C5-WT1-S01", 75, "Number patterns first."),
        ("C5-WT1-S03", "C5-WT1-S01, C5-WT1-S02", None, None),
        ("C5-FR-S01", None, None, None),  # not decided yet: left as it is
    ])
    plan = _run(db, path)
    assert (len(plan.add), plan.skills_left_alone) == (3, 1)
    assert _links(db) == [
        ("C5-WT1-S02", "C5-WT1-S01", 75), ("C5-WT1-S03", "C5-WT1-S01", None), ("C5-WT1-S03", "C5-WT1-S02", None),
    ]
    again = apply_prerequisites.make_plan(db, apply_prerequisites.read_rows(path))
    assert (again.add, again.change, again.remove, again.unchanged) == ([], [], [], 3)

    # S03 now needs only S02 at 80%; S02 needs nothing (None).
    _run(db, _sheet(tmp_path, [("C5-WT1-S03", "C5-WT1-S02", 80, None), ("C5-WT1-S02", "None", None, None)], "b.xlsx"))
    assert _links(db) == [("C5-WT1-S03", "C5-WT1-S02", 80)]
    entries = [json.loads(e.event_data_json) for e in db.query(AuditLog).order_by(AuditLog.created_at).all()]
    assert len(entries) == 2 and entries[1]["removed"] == [
        {"skill": "C5-WT1-S03", "needs": "C5-WT1-S01"}, {"skill": "C5-WT1-S02", "needs": "C5-WT1-S01"},
    ]


def test_anything_wrong_stops_the_run_before_writing(db, tmp_path):
    with pytest.raises(StopBeforeWriting) as stop:
        apply_prerequisites.read_rows(_sheet(tmp_path, [
            ("C5-WT1-S02", "C5-WT1-S02", None, None),
            ("C5-WT1-S03", "C5-WT1-S01", 150, None),
            ("C5-WT1-S03", "C5-WT1-S01", None, None),
        ]))
    assert str(stop.value).splitlines() == [
        "Row 2 (C5-WT1-S02): a skill can't need itself.",
        "Row 3 (C5-WT1-S03): Minimum Mastery % must be a whole number from 1 to 100, not 150.",
        "Row 4 (C5-WT1-S03): this skill appears more than once.",
    ]
    rows = apply_prerequisites.read_rows(_sheet(tmp_path, [("C5-WT1-S02", "C5-XX-S09", None, None)], "u.xlsx"))
    with pytest.raises(StopBeforeWriting, match="C5-XX-S09: no skill on the platform has this Skill ID."):
        apply_prerequisites.make_plan(db, rows)
    circle = apply_prerequisites.read_rows(_sheet(tmp_path, [
        ("C5-WT1-S01", "C5-WT1-S03", None, None),
        ("C5-WT1-S02", "C5-WT1-S01", None, None),
        ("C5-WT1-S03", "C5-WT1-S02", None, None),
    ], "c.xlsx"))
    with pytest.raises(StopBeforeWriting, match="in a circle"):
        apply_prerequisites.make_plan(db, circle)
    db.add(ConceptLesson(chapter_id="ch-3", code="C5-FR-S01", title="Same code", sequence=1))
    db.commit()
    same = apply_prerequisites.read_rows(_sheet(tmp_path, [("C5-WT1-S02", "C5-FR-S01", None, None)], "d.xlsx"))
    with pytest.raises(StopBeforeWriting, match="more than one skill has this Skill ID"):
        apply_prerequisites.make_plan(db, same)
    assert _links(db) == []


def test_the_sheet_is_made_from_the_chapter_workbooks(tmp_path):
    folder = tmp_path / "books"
    folder.mkdir()
    book = openpyxl.Workbook()
    bank = book.active
    bank.title = "Question Bank"
    bank.append(["Question ID", "Chapter No.", "Chapter", "Skill ID"])
    bank.append(["CBSE-C5-CH01-Q001", 1, "We the Travellers—I", "C5-WT1-S01"])
    skills = book.create_sheet("Skill Map")
    skills.append(["MathPath | Chapter 1 Skill Map"])
    skills.append([])
    skills.append(["Skill ID", "Skill", "Learning Outcome", "Prerequisite / Gap Check"])
    skills.append(["C5-WT1-S01", "Number patterns", "Continue sequences.", "Skip counting in 10s."])
    skills.append(["C5-WT1-S02", "Read numbers", "Read to 99,999.", "Number names."])
    skills.append(["Totals", None, None, None])
    book.save(folder / "CBSE_Class_5_Chapter_1.xlsx")
    out = tmp_path / "sheet.xlsx"
    make_workbook.write(make_workbook.read_skills(folder), out)
    rows = list(openpyxl.load_workbook(out)["Prerequisites"].iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS)
    assert rows[1][:5] == (1, "We the Travellers—I", "C5-WT1-S01", "Number patterns", "Skip counting in 10s.")
    assert len(rows) == 3
    # Read back as it is: every skill left empty, so nothing would change.
    read = apply_prerequisites.read_rows(str(out))
    assert [r.code for r in read] == ["C5-WT1-S01", "C5-WT1-S02"] and not any(r.needs or r.remove_all for r in read)


def test_sheet_order_is_kept_and_excel_percentages_are_read(db, tmp_path):
    import openpyxl as xl

    path = _sheet(tmp_path, [("C5-WT1-S03", "C5-WT1-S02, C5-WT1-S01", None, None)])
    book = xl.load_workbook(path)
    book["Prerequisites"].cell(row=2, column=HEADERS.index("Minimum Mastery %") + 1, value=0.75)
    book.save(path)
    plan = _run(db, path)
    assert [(s, n, m) for s, n, m, _t in plan.add] == [("C5-WT1-S03", "C5-WT1-S02", 75), ("C5-WT1-S03", "C5-WT1-S01", 75)]
    by_id = {l.id: l.code for l in db.query(ConceptLesson).all()}
    ordered = db.query(PrerequisiteLink).order_by(PrerequisiteLink.created_at).all()
    assert [by_id[l.prerequisite_concept_lesson_id] for l in ordered] == ["C5-WT1-S02", "C5-WT1-S01"]
    # Neither needed skill has a published practice set: said in the plan.
    assert plan.no_rescue_set == ["C5-WT1-S01", "C5-WT1-S02"]
