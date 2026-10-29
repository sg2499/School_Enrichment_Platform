"""School setup in one workbook, and sign-in slips (#210, 29 Oct 2026).

Rules under test: the template downloads; a workbook is checked whole and
nothing is written while any row has a problem (every problem listed by
sheet and row); a clean workbook creates teachers (with their mobile and
the sections they teach), students (with class, section and roll number)
and first passwords, shown once; uploading it again changes nothing; a
section already taught by someone else is refused; the sections list
counts each section's students and who teaches it; New Slips gives fresh
first passwords only to students who have not chosen their own; only the
school's admin (or the Super Admin, for a school they name) may do any of it.
"""

import io

from openpyxl import Workbook, load_workbook

from app.models import AuditLog, Student, Teacher, User
from app.models.teacher_assignment import TeacherSectionAssignment
from app.services import school_setup_service as setup
from tests.test_learning import (
    _get_or_create_board_course,
    _login,
    _make_school,
    _make_school_admin,
    _make_super_admin,
    _make_teacher,
)


def _xlsx(teachers=(), students=()) -> bytes:
    wb = Workbook()
    wb.active.title = "How to Fill In"
    for title, columns, rows in (
        ("Teachers", setup.TEACHER_COLUMNS, teachers),
        ("Students", setup.STUDENT_COLUMNS, students),
    ):
        ws = wb.create_sheet(title)
        ws.append(columns)
        for r in rows:
            ws.append(list(r))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _world(db, sx):
    """A school on a board of its own, so the one Class 5 Mathematics course
    is this test's (the suite shares one database, and other tests make
    CBSE Class 5 Mathematics courses of their own)."""
    from app.models.curriculum import Board, BoardCourse, ClassLevel

    _get_or_create_board_course(db, sx)  # makes sure Class 5 exists
    board = Board(code=f"B{sx}", display_name=f"Board {sx}")
    db.add(board)
    db.flush()
    level = db.query(ClassLevel).filter(ClassLevel.code == "5").one()
    course = BoardCourse(
        board_id=board.id, class_level_id=level.id, code=f"B{sx}-5-MATH", display_name="Mathematics", status="PUBLISHED"
    )
    db.add(course)
    school = _make_school(db, sx)
    school.board = board.code
    db.commit()
    admin = _make_school_admin(db, school, sx)
    return school, admin, course


def _student_login(code: str, password: str) -> int:
    """Signs in on a second client, so the admin's session in `client` is untouched."""
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as other:
        return other.post(
            "/api/auth/login", json={"identifier": code, "password": password, "signInAs": "STUDENT"}
        ).status_code


def _post(client, headers, path, data: bytes, school_id=None):
    form = {"schoolId": school_id} if school_id else {}
    return client.post(
        f"/api/roster/{path}",
        headers=headers,
        data=form,
        files={"file": ("setup.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


GOOD_TEACHERS = [
    ("Anita Rao", "anita.setup@example.com", "98765 43210", "", "5A, 5B"),
    ("Joseph Thomas", "", "+91 9123456780", "Maths", "5C"),
]
GOOD_STUDENTS = [
    ("Aarav Sharma", 5, "A", 1),
    ("Meera Iyer", 5, "A", 2),
    ("Kabir Khan", "5", "b", "1"),
    ("Example: Someone (delete this row)", 5, "A", 99),
]


def test_template_and_a_clean_import(client, db_session):
    db = db_session
    school, admin, course = _world(db, "SET1")
    headers = {**_login(client, admin.email), "x-auth-role": "ADMIN"}

    template = client.get("/api/roster/setup-template", headers=headers)
    assert template.status_code == 200
    wb = load_workbook(io.BytesIO(template.content))
    assert wb.sheetnames == ["How to Fill In", "Teachers", "Students"]
    assert [c.value for c in wb["Teachers"][1]] == setup.TEACHER_COLUMNS

    data = _xlsx(GOOD_TEACHERS, GOOD_STUDENTS)
    checked = _post(client, headers, "setup/check", data)
    assert checked.status_code == 200, checked.text
    body = checked.json()
    assert body["ok"] and body["problems"] == []
    assert (body["newTeachers"], body["newStudents"], body["newSectionsTaught"]) == (2, 3, 3)
    assert body["sections"] == ["5A", "5B"]
    assert db.query(Student).filter(Student.school_id == school.id).count() == 0  # the check wrote nothing

    imported = _post(client, headers, "setup/import", data)
    assert imported.status_code == 200, imported.text
    assert imported.headers["cache-control"] == "no-store"
    created = imported.json()["created"]
    assert [c["role"] for c in created] == ["TEACHER", "TEACHER", "STUDENT", "STUDENT", "STUDENT"]
    assert [(c["section"], c["rollNumber"]) for c in created if c["role"] == "STUDENT"] == [
        ("A", "1"),
        ("A", "2"),
        ("B", "1"),
    ]
    assert all(len(c["initialPassword"]) == 12 for c in created)
    anita = db.query(User).filter(User.email == "anita.setup@example.com").one()
    assert anita.phone == "9876543210" and anita.must_change_password
    joseph = (
        db.query(User)
        .join(Teacher, Teacher.user_id == User.id)
        .filter(User.full_name == "Joseph Thomas", Teacher.school_id == school.id)
        .one()
    )
    assert joseph.phone == "9123456780"
    teaches = {
        (a.section, a.board_course_id)
        for a in db.query(TeacherSectionAssignment)
        .join(Teacher, Teacher.id == TeacherSectionAssignment.teacher_id)
        .filter(Teacher.user_id == anita.id)
    }
    assert teaches == {("A", course.id), ("B", course.id)}
    kabir = (
        db.query(Student)
        .join(User, User.id == Student.user_id)
        .filter(User.full_name == "Kabir Khan", Student.school_id == school.id)
        .one()
    )
    assert (kabir.class_name, kabir.section, kabir.roll_number) == ("5", "B", "1") and kabir.class_level_id
    # A student can sign in with their code and the slip's password.
    slip = next(c for c in created if c["fullName"] == "Aarav Sharma")
    assert _student_login(slip["code"], slip["initialPassword"]) == 200
    assert db.query(AuditLog).filter(AuditLog.event_type == "roster.school_setup_import").count() == 1

    # The same file again: everyone is recognised, nothing is created.
    again = _post(client, headers, "setup/check", data).json()
    assert again["ok"] and (again["newTeachers"], again["newStudents"], again["newSectionsTaught"]) == (0, 0, 0)
    assert (again["teachersAlreadyHere"], again["studentsAlreadyHere"]) == (2, 3)
    assert _post(client, headers, "setup/import", data).json()["created"] == []
    # A corrected file with one more student and one more section adds just those.
    more = _xlsx(
        [*GOOD_TEACHERS[:1], ("Anita Rao", "anita.setup@example.com", "", "", "5D")],
        [*GOOD_STUDENTS, ("Rohan Das", 5, "A", 3)],
    )
    plus = _post(client, headers, "setup/import", more).json()
    assert [c["fullName"] for c in plus["created"]] == ["Rohan Das"] and plus["newSectionsTaught"] == 1

    sections = client.get("/api/roster/sections", headers=headers).json()["sections"]
    a = next(s for s in sections if s["section"] == "A")
    assert (a["className"], a["students"], a["awaitingFirstSignIn"]) == ("5", 3, 3)
    assert a["teachers"] == [{"fullName": "Anita Rao", "subject": "Mathematics"}]
    assert [s["section"] for s in sections] == ["A", "B", "C", "D"]


def test_every_problem_is_listed_and_nothing_is_written(client, db_session):
    db = db_session
    school, admin, course = _world(db, "SET2")
    other = _make_teacher(db, school, "set2-held")
    from tests.test_learning import _assign_teacher_section

    _assign_teacher_section(db, other, course.id, "A")
    elsewhere = _make_teacher(db, _make_school(db, "set2-else"), "set2-else")
    taken = elsewhere.user.email  # someone at another school
    headers = {**_login(client, admin.email), "x-auth-role": "ADMIN"}
    data = _xlsx(
        [
            ("", "", "", "", "5A"),
            ("Ravi", taken, "", "", "5B"),
            ("Sana", "", "12345", "", "5B"),
            ("Tenzin", "", "", "", "Fifth A"),
            ("Neha", "", "", "", "9A"),
            ("Divya", "", "", "Science", "5B"),
            ("Imran", "", "", "", "5A"),
            ("Priya", "", "", "", ""),
            ("Ravi Kumar", other.user.email, "", "", "5C"),
        ],
        [
            ("Aarav", 5, "A", 1),
            ("Meera", 5, "A", 1),
            ("Kabir", 11, "A", 2),
            ("Fatima", 5, "A1", 3),
            ("", 5, "A", 4),
        ],
    )
    body = _post(client, headers, "setup/check", data).json()
    assert not body["ok"]
    got = [(p["sheet"], p["row"]) for p in body["problems"]]
    # Row 3's address is someone's at another school: the check doesn't say (A14); the import does.
    assert got == [("Teachers", r) for r in (2, *range(4, 11))] + [
        ("Students", 3),
        ("Students", 4),
        ("Students", 5),
        ("Students", 6),
    ], body["problems"]
    words = " | ".join(p["message"] for p in body["problems"])
    for phrase in (
        "Full Name is empty",
        "10-digit",
        "isn't a section",
        "no Class 9",
        "no Science for BSET2 Class 5",
        "already taught by Teacher set2-held",
        "Sections is empty",
        "Roll number 1 in 5A is also on row 2",
        "isn't a class Krama has",
        "belongs to Teacher set2-held, who is already on Krama",
    ):
        assert phrase in words, phrase
    assert taken not in words  # the address is never echoed
    refused = _post(client, headers, "setup/import", data)
    assert refused.status_code == 422 and refused.json()["detail"]["code"] == "SETUP_HAS_PROBLEMS"
    at_import = refused.json()["detail"]["details"]["problems"]
    assert {"sheet": "Teachers", "row": 3, "message": setup.EMAIL_NOT_ACCEPTED} in at_import
    assert taken not in refused.text
    audit = db.query(AuditLog).filter(
        AuditLog.event_type == "roster.school_setup_import_refused", AuditLog.user_id == admin.id
    )
    assert audit.count() == 1 and taken not in str(audit.one().event_data_json)
    assert db.query(Student).filter(Student.school_id == school.id).count() == 0
    assert db.query(Teacher).filter(Teacher.school_id == school.id).count() == 1
    assert db.query(TeacherSectionAssignment).filter(TeacherSectionAssignment.school_id == school.id).count() == 1


def test_new_slips_only_for_students_who_have_not_signed_in(client, db_session):
    db = db_session
    school, admin, _course = _world(db, "SET3")
    headers = {**_login(client, admin.email), "x-auth-role": "ADMIN"}
    created = _post(
        client, headers, "setup/import", _xlsx([], [("Aarav Sharma", 5, "A", 1), ("Meera Iyer", 5, "A", 2)])
    ).json()["created"]
    meera = (
        db.query(User)
        .join(Student, Student.user_id == User.id)
        .filter(User.full_name == "Meera Iyer", Student.school_id == school.id)
        .one()
    )
    meera.must_change_password = False  # she has chosen her own
    db.commit()
    slips = client.post("/api/roster/sections/slips", headers=headers, json={"className": "5", "section": "a"})
    assert slips.status_code == 200, slips.text
    body = slips.json()
    assert [s["fullName"] for s in body["slips"]] == ["Aarav Sharma"] and body["alreadySignedIn"] == 1
    assert (
        client.post("/api/roster/sections/slips", headers=headers, json={"className": "5", "section": "Z"}).status_code
        == 404
    )
    old = next(c for c in created if c["fullName"] == "Aarav Sharma")
    assert body["slips"][0]["initialPassword"] != old["initialPassword"]
    assert _student_login(old["code"], old["initialPassword"]) == 401
    assert _student_login(old["code"], body["slips"][0]["initialPassword"]) == 200


def test_only_the_schools_admin(client, db_session):
    db = db_session
    school, admin, _course = _world(db, "SET4")
    other_school = _make_school(db, "set4-other")
    teacher = _make_teacher(db, school, "set4-t")
    data = _xlsx([], [("Aarav Sharma", 5, "A", 1)])
    t_headers = {**_login(client, teacher.user.email), "x-auth-role": "TEACHER"}
    assert _post(client, t_headers, "setup/check", data).status_code == 403
    assert client.get("/api/roster/sections", headers=t_headers).status_code == 403
    a_headers = {**_login(client, admin.email), "x-auth-role": "ADMIN"}
    assert _post(client, a_headers, "setup/check", data, school_id=other_school.id).status_code == 403
    assert _post(client, a_headers, "setup/check", b"not a workbook").status_code == 422
    sa = _make_super_admin(db, "set4")
    s_headers = {**_login(client, sa.email), "x-auth-role": "SUPER_ADMIN"}
    assert _post(client, s_headers, "setup/check", data).status_code == 422  # must choose a school
    assert _post(client, s_headers, "setup/check", data, school_id=school.id).json()["newStudents"] == 1


def test_review_fixes(client, db_session):
    """The 29 Oct review's findings, each as it was found."""
    import time

    db = db_session
    school, admin, _course = _world(db, "SET5")
    headers = {**_login(client, admin.email), "x-auth-role": "ADMIN"}
    elsewhere = _make_teacher(db, _make_school(db, "set5-else"), "set5-else")
    elsewhere.user.phone = "9000012345"
    db.commit()

    # 1. The check never says whether an email or mobile is on Krama elsewhere.
    probe = _xlsx([("P One", elsewhere.user.email, "", "", "5A"), ("P Two", "", "9000012345", "", "5B")])
    assert _post(client, headers, "setup/check", probe).json()["ok"] is True
    words = [
        p["message"] for p in _post(client, headers, "setup/import", probe).json()["detail"]["details"]["problems"]
    ]
    assert words == [setup.EMAIL_NOT_ACCEPTED, setup.MOBILE_NOT_ACCEPTED]

    # 4. Two different teachers can't share an email or a mobile.
    shared = _xlsx(
        [
            ("Anita Rao", "x5@example.com", "", "", "5A"),
            ("Ravi Kumar", "x5@example.com", "", "", "5B"),
            ("Sunil Das", "", "9811122233", "", "5C"),
            ("Gita Das", "", "98111 22233", "", "5D"),
        ]
    )
    body = _post(client, headers, "setup/check", shared).json()
    assert [
        (p["row"], "for Anita Rao" in p["message"] or "for Sunil Das" in p["message"]) for p in body["problems"]
    ] == [
        (3, True),
        (5, True),
    ]
    # ...but a teacher's further subject row needn't repeat the contact details.
    two_rows = _xlsx([("Anita Rao", "x5@example.com", "", "", "5A"), ("Anita Rao", "", "", "Mathematics", "5B")])
    body = _post(client, headers, "setup/check", two_rows).json()
    assert body["ok"] and body["newTeachers"] == 1 and body["newSectionsTaught"] == 2

    # 5. One child already here without a roll number can't absorb two rows of that name.
    assert _post(client, headers, "setup/import", _xlsx([], [("Aarav Shah", 5, "A", "")])).status_code == 200
    body = _post(
        client, headers, "setup/check", _xlsx([], [("Aarav Shah", 5, "A", 1), ("Aarav Shah", 5, "A", 2)])
    ).json()
    assert [p["row"] for p in body["problems"]] == [3] and "can't tell which row" in body["problems"][0]["message"]

    # 6. Same name twice in a file: both need a roll number, so the file uploads again cleanly.
    body = _post(
        client, headers, "setup/check", _xlsx([], [("Diya Nair", 5, "B", 1), ("Diya Nair", 5, "B", "")])
    ).json()
    assert [p["row"] for p in body["problems"]] == [3]
    twins = _xlsx([], [("Diya Nair", 5, "B", 1), ("Diya Nair", 5, "B", 2)])
    assert _post(client, headers, "setup/import", twins).json()["newStudents"] == 2
    again = _post(client, headers, "setup/check", twins).json()
    assert again["ok"] and again["newStudents"] == 0 and again["studentsAlreadyHere"] == 2

    # 12. Roll "01", then 1 once Excel drops the zero: the same child.
    assert _post(client, headers, "setup/import", _xlsx([], [("Rohan Das", 5, "C", "01")])).status_code == 200
    again = _post(client, headers, "setup/check", _xlsx([], [("Rohan Das", 5, "C", 1)])).json()
    assert again["ok"] and again["studentsAlreadyHere"] == 1

    # 2. New Slips for a class People named freely ("UKG") works, and is on each child's own trail.
    from tests.test_learning import _make_student

    ukg = _make_student(db, school, "set5-ukg")
    ukg.class_name, ukg.section = "UKG", "A"
    ukg.user.must_change_password = True
    db.commit()
    slips = client.post("/api/roster/sections/slips", headers=headers, json={"className": "UKG", "section": "A"})
    assert slips.status_code == 200, slips.text
    assert _student_login(ukg.student_code, slips.json()["slips"][0]["initialPassword"]) == 200
    trail = db.query(AuditLog).filter(
        AuditLog.event_type == "auth.password_reset_by_admin", AuditLog.user_id == ukg.user_id
    )
    assert trail.count() == 1

    # 3. A sheet with a vast number of rows is refused without reading them all.
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(setup.STUDENT_COLUMNS)
    for _ in range(25_000):
        ws.append([None, None, None, " "])  # blank to the eye, but each row is still there to read
    out = io.BytesIO()
    wb.save(out)
    start = time.monotonic()
    refused = _post(client, headers, "setup/check", out.getvalue())
    assert refused.status_code == 422 and "row" in refused.json()["detail"]["message"]
    assert time.monotonic() - start < 5
